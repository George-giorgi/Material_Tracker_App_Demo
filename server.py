import csv
import json
import hashlib
import hmac
import math
import os
import re
import secrets
import sqlite3
import ssl
import sys
import time
from threading import Lock
from urllib.error import URLError
from urllib.request import Request, urlopen
import zipfile
from contextlib import contextmanager
from datetime import date, datetime
from http.cookies import CookieError, SimpleCookie
from io import BytesIO, StringIO
from xml.etree.ElementTree import ParseError, fromstring
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse
from xml.sax.saxutils import escape

import pymupdf as fitz
from pypdf.errors import PdfReadError
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
import reportlab


ROOT = Path(__file__).resolve().parent
STATIC_DIR = ROOT / "static"
DB_PATH = Path(os.environ.get("EWI_DB_PATH", ROOT / "instance" / "ewi.sqlite3"))
PORT = int(os.environ.get("EWI_PORT", "8001"))
HOST = os.environ.get("EWI_HOST", "127.0.0.1")
TLS_CERT = os.environ.get("EWI_TLS_CERT", "")
TLS_KEY = os.environ.get("EWI_TLS_KEY", "")
COOKIE_SECURE = bool(TLS_CERT and TLS_KEY)
SESSION_MAX_AGE = 12 * 60 * 60
CREDENTIAL_ITERATIONS = 310_000
HOUSE_DOCUMENT_MAX_BYTES = 8_000_000
HOUSE_DOCUMENT_MAX_COUNT = 25
HOUSE_DOCUMENT_MAX_PAGES = 300
HOUSE_DOCUMENT_MAX_ANNOTATIONS = 500
HOUSE_DOCUMENT_MAX_ANNOTATION_POINTS = 1000
SEAI_DOCUMENT_MAX_BYTES = 20_000_000
MOVEMENT_TYPES = {"collection", "pickup", "delivery", "use", "return", "transfer"}
LOCATION_TYPES = {"warehouse", "team", "house"}
MET_EIREANN_COUNTY_FORECAST_URL = (
    "https://www.met.ie/Open_Data/xml/county_forecast.xml"
)
MET_EIREANN_WEATHER_CACHE_SECONDS = 15 * 60
_met_eireann_weather_cache = None
_met_eireann_weather_cache_expiry = 0
_met_eireann_weather_lock = Lock()


class ApiError(Exception):
    def __init__(self, status, message):
        super().__init__(message)
        self.status = status


def _parse_met_eireann_county_forecast(content):
    try:
        root = fromstring(content)
    except ParseError as exc:
        raise ApiError(502, "Met Éireann returned an unreadable forecast.") from exc
    if root.tag != "forecast":
        raise ApiError(502, "Met Éireann returned an unexpected forecast format.")

    wanted_counties = {"MAYO": "Mayo", "SLIGO": "Sligo"}
    areas = []
    for county in root.findall("county"):
        county_name = (county.findtext("name") or "").strip().upper()
        if county_name not in wanted_counties:
            continue
        days = []
        for forecast_day in county.findall("day"):
            forecast_date = (forecast_day.findtext("date") or "").strip()
            try:
                datetime.fromisoformat(forecast_date)
            except ValueError as exc:
                raise ApiError(
                    502, "Met Éireann returned an invalid forecast date."
                ) from exc

            def number(field):
                raw_value = forecast_day.findtext(field)
                if raw_value is None or not raw_value.strip():
                    return None
                try:
                    value = float(raw_value)
                except ValueError as exc:
                    raise ApiError(
                        502, "Met Éireann returned invalid forecast values."
                    ) from exc
                return value if math.isfinite(value) else None

            rainfall_day = number("rainfall_6_18")
            rainfall_night = number("rainfall_18_6")
            days.append(
                {
                    "date": forecast_date[:10],
                    "weather": (forecast_day.findtext("weather") or "").strip(),
                    "min_temp": number("min_temp"),
                    "max_temp": number("max_temp"),
                    "wind_speed_kmh": number("wind_speed"),
                    "wind_direction": (forecast_day.findtext("wind_dir") or "").strip(),
                    "rainfall_mm": (
                        rainfall_day + rainfall_night
                        if rainfall_day is not None and rainfall_night is not None
                        else None
                    ),
                }
            )
        if days:
            areas.append(
                {"county": wanted_counties[county_name], "days": days}
            )

    if {area["county"] for area in areas} != {"Mayo", "Sligo"}:
        raise ApiError(
            502, "The Mayo and Sligo county forecasts are temporarily unavailable."
        )
    return {
        "issued": root.attrib.get("issued"),
        "areas": areas,
        "source": "Met Éireann Open Data",
        "source_url": "https://www.met.ie/",
    }


def _met_eireann_county_forecast():
    global _met_eireann_weather_cache, _met_eireann_weather_cache_expiry
    with _met_eireann_weather_lock:
        now = time.monotonic()
        if (
            _met_eireann_weather_cache is not None
            and now < _met_eireann_weather_cache_expiry
        ):
            return _met_eireann_weather_cache

        request = Request(
            MET_EIREANN_COUNTY_FORECAST_URL,
            headers={"User-Agent": "MaterialTracker/1.0"},
        )
        try:
            with urlopen(request, timeout=8) as response:
                content = response.read(256_001)
        except (URLError, TimeoutError, OSError) as exc:
            raise ApiError(
                502, "Met Éireann weather is temporarily unavailable."
            ) from exc
        if len(content) > 256_000:
            raise ApiError(502, "Met Éireann returned an oversized forecast.")

        forecast = _parse_met_eireann_county_forecast(content)
        _met_eireann_weather_cache = forecast
        _met_eireann_weather_cache_expiry = (
            time.monotonic() + MET_EIREANN_WEATHER_CACHE_SECONDS
        )
        return forecast


def _open_pdf_document(content):
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except (fitz.FileDataError, ValueError, RuntimeError) as exc:
        raise ApiError(422, "This uploaded file could not be read as a PDF.") from exc
    if document.needs_pass:
        document.close()
        raise ApiError(422, "Password-protected PDFs cannot be annotated.")
    if not document.page_count:
        document.close()
        raise ApiError(422, "This PDF does not contain any pages.")
    if document.page_count > HOUSE_DOCUMENT_MAX_PAGES:
        document.close()
        raise ApiError(
            422,
            f"PDFs with more than {HOUSE_DOCUMENT_MAX_PAGES} pages cannot be annotated.",
        )
    return document


def _render_pdf_page(content, page_number):
    document = _open_pdf_document(content)
    try:
        if page_number < 1 or page_number > document.page_count:
            raise ApiError(404, "That PDF page does not exist.")
        page = document.load_page(page_number - 1)
        rect = page.rect
        scale = min(1.6, 1800 / max(rect.width, rect.height))
        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        return pixmap.tobytes("png"), document.page_count
    except ApiError:
        raise
    except (fitz.FileDataError, ValueError, RuntimeError) as exc:
        raise ApiError(422, "This PDF page could not be rendered.") from exc
    finally:
        document.close()


def _annotate_pdf_document(content, annotations):
    if not isinstance(annotations, list) or not annotations:
        raise ApiError(400, "Add at least one pen stroke or text note before saving.")
    if len(annotations) > HOUSE_DOCUMENT_MAX_ANNOTATIONS:
        raise ApiError(400, "This PDF has too many annotations to save at once.")
    document = _open_pdf_document(content)
    try:
        for annotation in annotations:
            if not isinstance(annotation, dict):
                raise ApiError(400, "PDF annotations are invalid.")
            page_number = annotation.get("page")
            tool = annotation.get("tool")
            color_value = annotation.get("color")
            if (
                isinstance(page_number, bool)
                or not isinstance(page_number, int)
                or not 1 <= page_number <= document.page_count
                or tool not in {"pen", "highlighter", "text"}
                or not isinstance(color_value, str)
                or not re.fullmatch(r"#[0-9a-fA-F]{6}", color_value)
            ):
                raise ApiError(400, "PDF annotations contain an invalid page, tool, or color.")
            page = document.load_page(page_number - 1)
            rect = page.rect
            color = tuple(
                int(color_value[index:index + 2], 16) / 255
                for index in (1, 3, 5)
            )
            if tool == "text":
                text = annotation.get("text")
                point = annotation.get("point")
                font_ratio = annotation.get("font_size")
                if (
                    not isinstance(text, str)
                    or not text.strip()
                    or len(text) > 240
                    or not isinstance(point, list)
                    or len(point) != 2
                    or not isinstance(font_ratio, (int, float))
                    or not 0.005 <= font_ratio <= 0.1
                    or not all(
                        isinstance(value, (int, float)) and 0 <= value <= 1
                        for value in point
                    )
                ):
                    raise ApiError(400, "A text annotation is invalid.")
                font_size = rect.width * font_ratio
                page.insert_text(
                    (
                        point[0] * rect.width,
                        min(rect.height, point[1] * rect.height + font_size),
                    ),
                    text,
                    fontsize=font_size,
                    fontname="helv",
                    color=color,
                    overlay=True,
                )
                continue
            points = annotation.get("points")
            width_ratio = annotation.get("width")
            if (
                not isinstance(points, list)
                or not points
                or len(points) > HOUSE_DOCUMENT_MAX_ANNOTATION_POINTS
                or not isinstance(width_ratio, (int, float))
                or not 0.0005 <= width_ratio <= 0.05
            ):
                raise ApiError(400, "A pen annotation is invalid.")
            normalized_points = []
            for point in points:
                if (
                    not isinstance(point, list)
                    or len(point) != 2
                    or not all(
                        isinstance(value, (int, float)) and 0 <= value <= 1
                        for value in point
                    )
                ):
                    raise ApiError(400, "A pen annotation is outside the PDF page.")
                normalized_points.append(
                    fitz.Point(point[0] * rect.width, point[1] * rect.height)
                )
            width = max(0.5, rect.width * width_ratio)
            if len(normalized_points) == 1:
                point = normalized_points[0]
                page.draw_circle(
                    point,
                    width / 2,
                    color=color,
                    fill=color,
                    width=0,
                    overlay=True,
                )
            else:
                page.draw_polyline(
                    normalized_points,
                    color=color,
                    width=width,
                    lineCap=1,
                    lineJoin=1,
                    stroke_opacity=0.38 if tool == "highlighter" else 1,
                    overlay=True,
                )
        return document.tobytes(garbage=4, deflate=True)
    except ApiError:
        raise
    except (fitz.FileDataError, ValueError, RuntimeError) as exc:
        raise ApiError(422, "The annotations could not be added to this PDF.") from exc
    finally:
        document.close()


def _create_house_document_share_pdf(content, filename, house_id, survey_note):
    try:
        original = PdfReader(BytesIO(content), strict=False)
    except (PdfReadError, ValueError, OSError) as exc:
        raise ApiError(422, "This uploaded file could not be read as a PDF.") from exc
    if original.is_encrypted:
        raise ApiError(422, "Password-protected PDFs cannot be prepared for sharing.")
    if not original.pages:
        raise ApiError(422, "This PDF does not contain any pages.")
    if len(original.pages) > HOUSE_DOCUMENT_MAX_PAGES:
        raise ApiError(
            422,
            f"PDFs with more than {HOUSE_DOCUMENT_MAX_PAGES} pages cannot be prepared for sharing.",
        )

    font_path = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
    font_name = "TrackerSurveyNotes"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "SurveyNoteTitle",
        parent=styles["Title"],
        fontName="Times-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#233247"),
        alignment=TA_LEFT,
        spaceAfter=6 * mm,
    )
    details_style = ParagraphStyle(
        "SurveyNoteDetails",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#303942"),
    )
    details_label_style = ParagraphStyle(
        "SurveyNoteDetailsLabel",
        parent=details_style,
        fontSize=7,
        leading=10,
        textColor=colors.HexColor("#68727c"),
        tracking=0.6,
    )
    note_style = ParagraphStyle(
        "SurveyNoteBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=10.5,
        leading=16,
        textColor=colors.HexColor("#303942"),
        leftIndent=13,
        bulletIndent=0,
        bulletFontName=font_name,
        bulletFontSize=10,
        bulletColor=colors.HexColor("#34465a"),
    )
    note_lines = [line.strip() for line in survey_note.splitlines() if line.strip()]
    if not note_lines:
        note_lines = [survey_note.strip()]
    note_pdf = BytesIO()
    note_document = SimpleDocTemplate(
        note_pdf,
        pagesize=A4,
        rightMargin=22 * mm,
        leftMargin=22 * mm,
        topMargin=30 * mm,
        bottomMargin=24 * mm,
        title=f"Survey note - {house_id}",
    )

    def draw_page_chrome(pdf_canvas, document):
        page_width, page_height = A4
        pdf_canvas.saveState()
        pdf_canvas.setFillColor(colors.HexColor("#34465a"))
        pdf_canvas.setFont(font_name, 8)
        pdf_canvas.drawString(22 * mm, page_height - 17 * mm, "BRUSNA  |  MATERIAL TRACKER")
        pdf_canvas.setStrokeColor(colors.HexColor("#9aa5b1"))
        pdf_canvas.setLineWidth(0.7)
        pdf_canvas.line(22 * mm, page_height - 22 * mm, page_width - 22 * mm, page_height - 22 * mm)
        pdf_canvas.line(22 * mm, 17 * mm, page_width - 22 * mm, 17 * mm)
        pdf_canvas.setFillColor(colors.HexColor("#5f6973"))
        pdf_canvas.setFont(font_name, 7)
        pdf_canvas.drawString(22 * mm, 12 * mm, f"SURVEY NOTE  |  {house_id}")
        pdf_canvas.drawRightString(page_width - 22 * mm, 12 * mm, f"PAGE {document.page}")
        pdf_canvas.restoreState()

    metadata_table = Table(
        [
            [
                Paragraph("HOUSE", details_label_style),
                Paragraph(escape(house_id), details_style),
            ],
            [
                Paragraph("SOURCE PDF", details_label_style),
                Paragraph(escape(filename), details_style),
            ],
        ],
        colWidths=[34 * mm, 128 * mm],
        hAlign="LEFT",
    )
    metadata_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#d5dbe1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    note_document.build(
        [
            Paragraph("House survey note", title_style),
            metadata_table,
            Spacer(1, 9 * mm),
            Paragraph("SURVEY OBSERVATIONS", details_label_style),
            Spacer(1, 3 * mm),
            *[
                Paragraph(escape(line), note_style, bulletText="•")
                for line in note_lines
            ],
        ],
        onFirstPage=draw_page_chrome,
        onLaterPages=draw_page_chrome,
    )

    writer = PdfWriter()
    for page in original.pages:
        writer.add_page(page)
    note_reader = PdfReader(note_pdf)
    for page in note_reader.pages:
        writer.add_page(page)
    writer.add_metadata(
        {
            "/Title": f"{house_id} - {Path(filename).stem} with survey note",
            "/Subject": "Original survey PDF with the survey note appended on a separate page",
        }
    )
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _create_seai_document_share_pdf(selections, documents, house_id, survey_note):
    writer = PdfWriter()
    filenames = set()
    for selection in selections:
        source = documents[selection["document_id"]]
        try:
            reader = PdfReader(BytesIO(source["content"]), strict=False)
        except (PdfReadError, ValueError, OSError) as exc:
            raise ApiError(422, "A selected SEAI PDF could not be read.") from exc
        if reader.is_encrypted:
            raise ApiError(422, "Password-protected SEAI PDFs cannot be shared.")
        page_number = selection["page"]
        if page_number < 1 or page_number > len(reader.pages):
            raise ApiError(
                400,
                f"Page {page_number} does not exist in {source['filename']}.",
            )
        writer.add_page(reader.pages[page_number - 1])
        filenames.add(source["filename"])
    selected_pages = BytesIO()
    writer.write(selected_pages)
    source_label = (
        f"Selected SEAI pages ({len(selections)} pages from {len(filenames)} PDFs)"
    )
    return _create_house_document_share_pdf(
        selected_pages.getvalue(), source_label, house_id, survey_note
    )


def _seai_document_source_stem(filename):
    stem = Path(filename).stem
    return re.sub(
        r"(?:-annotated(?:-\d+)?)+$",
        "",
        stem,
        flags=re.IGNORECASE,
    )


def _is_annotated_pdf_filename(filename):
    return bool(
        re.search(r"(?:-annotated)+(?:-\d+)?\.pdf$", str(filename), re.IGNORECASE)
    )


def _is_not_started_house(db, house_id):
    row = db.execute(
        """
        SELECT p.ready_for_start, h.id AS active_id
        FROM provided_houses p
        LEFT JOIN houses h ON h.id = p.id
        WHERE p.id = ?
        """,
        (house_id,),
    ).fetchone()
    return bool(row and not row["ready_for_start"] and row["active_id"] is None)


def _require_pdf_editable(db, house_id):
    if _is_not_started_house(db, house_id):
        raise ApiError(409, "PDFs are view-only while the house is Not started.")


def _clear_house_survey_records(db, house_id):
    db.execute("DELETE FROM house_surveys WHERE house_id = ?", (house_id,))
    db.execute("DELETE FROM house_snags WHERE house_id = ?", (house_id,))
    db.execute("DELETE FROM house_checklists WHERE house_id = ?", (house_id,))
    db.execute(
        """
        UPDATE provided_houses
        SET notes = '',
            plan_scaffolding = NULL,
            plan_esb = NULL,
            plan_plumber = NULL,
            plan_scaffolding_required = NULL,
            plan_scaffolding_notes = '',
            plan_esb_required = NULL,
            plan_esb_notes = '',
            plan_plumber_required = NULL,
            plan_plumber_notes = ''
        WHERE id = ?
        """,
        (house_id,),
    )

    house_document_ids = [
        row["id"]
        for row in db.execute(
            "SELECT id, filename FROM house_documents WHERE house_id = ?",
            (house_id,),
        )
        if _is_annotated_pdf_filename(row["filename"])
    ]
    if house_document_ids:
        db.executemany(
            "DELETE FROM house_documents WHERE id = ?",
            ((document_id,) for document_id in house_document_ids),
        )

    seai_documents_table = db.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'seai_documents'"
    ).fetchone()
    if seai_documents_table:
        annotated_ids = [
            row["id"]
            for row in db.execute(
                """
                SELECT id, filename FROM seai_documents
                WHERE application_id = ?
                """,
                (house_id,),
            )
            if _is_annotated_pdf_filename(row["filename"])
        ]
        if annotated_ids:
            db.executemany(
                "DELETE FROM seai_documents WHERE id = ?",
                ((document_id,) for document_id in annotated_ids),
            )

@contextmanager
def connect_db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys = ON")
    try:
        yield db
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def initialize_db():
    with connect_db() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS materials (
                id INTEGER PRIMARY KEY,
                code TEXT,
                name TEXT NOT NULL COLLATE NOCASE UNIQUE,
                unit TEXT NOT NULL DEFAULT 'units',
                opening_qty REAL
            );
            CREATE TABLE IF NOT EXISTS houses (
                id TEXT PRIMARY KEY,
                postcode TEXT NOT NULL DEFAULT '',
                owner_name TEXT NOT NULL DEFAULT '',
                employee_name TEXT NOT NULL DEFAULT '',
                is_finished INTEGER NOT NULL DEFAULT 0,
                status_changed_at REAL NOT NULL DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS house_status_history (
                id INTEGER PRIMARY KEY,
                house_id TEXT NOT NULL,
                status TEXT NOT NULL,
                changed_by TEXT NOT NULL DEFAULT '',
                changed_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS house_status_history_house_idx
                ON house_status_history(house_id, id);
            CREATE TABLE IF NOT EXISTS house_checklists (
                house_id TEXT PRIMARY KEY REFERENCES provided_houses(id) ON DELETE CASCADE,
                scaffolding_status TEXT,
                esb_unclipped TEXT,
                plumber_status TEXT
            );
            CREATE TABLE IF NOT EXISTS movements (
                id INTEGER PRIMARY KEY,
                happened_on TEXT NOT NULL,
                kind TEXT NOT NULL,
                material_id INTEGER NOT NULL REFERENCES materials(id),
                qty REAL NOT NULL CHECK (qty > 0),
                source TEXT NOT NULL,
                destination TEXT NOT NULL,
                house_id TEXT REFERENCES houses(id),
                delivery_number INTEGER,
                person TEXT NOT NULL DEFAULT '',
                team TEXT NOT NULL DEFAULT '',
                department TEXT NOT NULL DEFAULT '',
                notes TEXT NOT NULL DEFAULT '',
                recorded_at TEXT NOT NULL DEFAULT '',
                FOREIGN KEY (material_id) REFERENCES materials(id)
            );
            CREATE INDEX IF NOT EXISTS movements_material_idx ON movements(material_id, id);
            CREATE TABLE IF NOT EXISTS stocktakes (
                id INTEGER PRIMARY KEY,
                material_id INTEGER NOT NULL REFERENCES materials(id),
                location TEXT NOT NULL,
                quantity REAL NOT NULL CHECK (quantity >= 0),
                movement_id INTEGER NOT NULL,
                counted_on TEXT NOT NULL,
                counted_by TEXT NOT NULL DEFAULT '',
                notes TEXT NOT NULL DEFAULT ''
            );
            CREATE INDEX IF NOT EXISTS stocktakes_location_idx
                ON stocktakes(material_id, location, id);
            CREATE TABLE IF NOT EXISTS app_settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS app_users (
                id INTEGER PRIMARY KEY,
                username TEXT NOT NULL UNIQUE COLLATE NOCASE,
                display_name TEXT NOT NULL,
                role TEXT NOT NULL CHECK (role IN ('admin', 'employee', 'surveyor')),
                password_salt BLOB NOT NULL,
                password_hash BLOB NOT NULL,
                password_iterations INTEGER NOT NULL,
                credential_type TEXT NOT NULL DEFAULT 'pin',
                is_active INTEGER NOT NULL DEFAULT 1
            );
            CREATE TABLE IF NOT EXISTS app_sessions (
                token_hash TEXT PRIMARY KEY,
                user_id INTEGER NOT NULL REFERENCES app_users(id) ON DELETE CASCADE,
                expires_at INTEGER NOT NULL
            );
            CREATE INDEX IF NOT EXISTS app_sessions_expiry_idx
                ON app_sessions(expires_at);
            CREATE TABLE IF NOT EXISTS app_login_attempts (
                username_hash TEXT PRIMARY KEY,
                failures INTEGER NOT NULL,
                window_started_at INTEGER NOT NULL,
                locked_until INTEGER NOT NULL DEFAULT 0
            );
            """
        )
        if db.execute(
            "SELECT 1 FROM sqlite_master WHERE type = 'table' "
            "AND name = 'seai_documents'"
        ).fetchone():
            seai_document_columns = {
                row["name"] for row in db.execute("PRAGMA table_info(seai_documents)")
            }
            if "content" not in seai_document_columns:
                db.execute("ALTER TABLE seai_documents ADD COLUMN content BLOB")
        status_history_foreign_keys = {
            row["table"]
            for row in db.execute("PRAGMA foreign_key_list(house_status_history)")
        }
        if "houses" in status_history_foreign_keys:
            db.execute(
                "ALTER TABLE house_status_history RENAME TO house_status_history_old"
            )
            db.execute(
                """
                CREATE TABLE house_status_history (
                    id INTEGER PRIMARY KEY,
                    house_id TEXT NOT NULL,
                    status TEXT NOT NULL,
                    changed_by TEXT NOT NULL DEFAULT '',
                    changed_at TEXT NOT NULL
                )
                """
            )
            db.execute(
                """
                INSERT INTO house_status_history(id, house_id, status, changed_by, changed_at)
                SELECT id, house_id, status, changed_by, changed_at
                FROM house_status_history_old
                """
            )
            db.execute("DROP TABLE house_status_history_old")
            db.execute(
                """
                CREATE INDEX house_status_history_house_idx
                ON house_status_history(house_id, id)
                """
            )
        house_columns = {
            row["name"] for row in db.execute("PRAGMA table_info(houses)")
        }
        if "is_archived" in house_columns and sqlite3.sqlite_version_info >= (3, 35, 0):
            db.execute("ALTER TABLE houses DROP COLUMN is_archived")
            house_columns.remove("is_archived")
        if "is_finished" not in house_columns:
            db.execute(
                "ALTER TABLE houses ADD COLUMN is_finished INTEGER NOT NULL DEFAULT 0"
            )
        if "status_changed_at" not in house_columns:
            db.execute(
                "ALTER TABLE houses ADD COLUMN status_changed_at REAL NOT NULL DEFAULT 0"
            )
        movement_columns = {
            row["name"] for row in db.execute("PRAGMA table_info(movements)")
        }
        if "house_id" not in movement_columns:
            db.execute("ALTER TABLE movements ADD COLUMN house_id TEXT REFERENCES houses(id)")
            db.execute(
                """
                UPDATE movements
                SET house_id = CASE
                    WHEN destination LIKE 'house:%' THEN substr(destination, 7)
                    WHEN source LIKE 'house:%' THEN substr(source, 7)
                    ELSE NULL
                END
                WHERE house_id IS NULL
                """
            )
        if "delivery_number" not in movement_columns:
            db.execute("ALTER TABLE movements ADD COLUMN delivery_number INTEGER")
        delivery_rows = db.execute(
            "SELECT id, house_id FROM movements "
            "WHERE kind = 'delivery' AND delivery_number IS NULL ORDER BY id"
        ).fetchall()
        delivery_numbers = {
            row["house_id"]: row["number"]
            for row in db.execute(
                """
                SELECT house_id, MAX(delivery_number) AS number
                FROM movements
                WHERE kind = 'delivery' AND delivery_number IS NOT NULL
                GROUP BY house_id
                """
            )
        }
        for row in delivery_rows:
            if not row["house_id"]:
                continue
            number = delivery_numbers.get(row["house_id"], 0) + 1
            delivery_numbers[row["house_id"]] = number
            db.execute(
                "UPDATE movements SET delivery_number = ? WHERE id = ?",
                (number, row["id"]),
            )
        house_columns = {row["name"] for row in db.execute("PRAGMA table_info(houses)")}
        if "employee_name" not in house_columns:
            db.execute(
                "ALTER TABLE houses ADD COLUMN employee_name TEXT NOT NULL DEFAULT ''"
            )
        if "stage" not in house_columns:
            db.execute("ALTER TABLE houses ADD COLUMN stage INTEGER NOT NULL DEFAULT 1")
        db.execute("UPDATE houses SET stage = 1 WHERE stage NOT IN (1, 2)")
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS provided_houses (
                id TEXT PRIMARY KEY,
                postcode TEXT NOT NULL DEFAULT '',
                address TEXT NOT NULL DEFAULT '',
                owner_name TEXT NOT NULL DEFAULT '',
                area_m2 REAL,
                perimeter_m REAL,
                source TEXT NOT NULL DEFAULT 'tracker'
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS house_snags (
                id INTEGER PRIMARY KEY,
                house_id TEXT NOT NULL REFERENCES provided_houses(id) ON DELETE CASCADE,
                text TEXT NOT NULL,
                done INTEGER NOT NULL DEFAULT 0,
                created_at REAL NOT NULL DEFAULT 0
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS house_surveys (
                house_id TEXT PRIMARY KEY REFERENCES provided_houses(id) ON DELETE CASCADE,
                data TEXT NOT NULL,
                notes TEXT NOT NULL DEFAULT '',
                updated_by TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS house_documents (
                id INTEGER PRIMARY KEY,
                house_id TEXT NOT NULL REFERENCES provided_houses(id) ON DELETE CASCADE,
                filename TEXT NOT NULL,
                content BLOB NOT NULL,
                size_bytes INTEGER NOT NULL,
                uploaded_by TEXT NOT NULL DEFAULT '',
                uploaded_at TEXT NOT NULL
            )
            """
        )
        provided_columns = {
            row["name"] for row in db.execute("PRAGMA table_info(provided_houses)")
        }
        for column, definition in (
            ("ready_for_start", "INTEGER NOT NULL DEFAULT 0"),
            ("notes", "TEXT NOT NULL DEFAULT ''"),
            ("plan_scaffolding", "TEXT"),
            ("plan_esb", "TEXT"),
            ("plan_plumber", "TEXT"),
            ("plan_scaffolding_required", "INTEGER"),
            ("plan_scaffolding_notes", "TEXT NOT NULL DEFAULT ''"),
            ("plan_esb_required", "INTEGER"),
            ("plan_esb_notes", "TEXT NOT NULL DEFAULT ''"),
            ("plan_plumber_required", "INTEGER"),
            ("plan_plumber_notes", "TEXT NOT NULL DEFAULT ''"),
            ("latitude", "REAL"),
            ("longitude", "REAL"),
        ):
            if column not in provided_columns:
                db.execute(f"ALTER TABLE provided_houses ADD COLUMN {column} {definition}")
        # Every house the tracker already works on belongs to the provided list.
        db.execute(
            """
            INSERT OR IGNORE INTO provided_houses(id, postcode, owner_name, source)
            SELECT id, postcode, owner_name, 'tracker' FROM houses
            """
        )
        checklist_parent = next(
            (
                row["table"]
                for row in db.execute("PRAGMA foreign_key_list(house_checklists)")
                if row["from"] == "house_id"
            ),
            None,
        )
        if checklist_parent != "provided_houses":
            db.execute("ALTER TABLE house_checklists RENAME TO house_checklists_old")
            db.execute(
                """
                CREATE TABLE house_checklists (
                    house_id TEXT PRIMARY KEY REFERENCES provided_houses(id) ON DELETE CASCADE,
                    scaffolding_status TEXT,
                    esb_unclipped TEXT,
                    plumber_status TEXT
                )
                """
            )
            db.execute(
                """
                INSERT INTO house_checklists
                    (house_id, scaffolding_status, esb_unclipped, plumber_status)
                SELECT house_id, scaffolding_status, esb_unclipped, plumber_status
                FROM house_checklists_old
                """
            )
            db.execute("DROP TABLE house_checklists_old")
        references_houses = any(
            row["table"] == "houses"
            for row in db.execute("PRAGMA foreign_key_list(material_requests)")
        )
        if references_houses:
            db.execute("ALTER TABLE material_requests RENAME TO material_requests_old")
            db.execute("DROP INDEX IF EXISTS material_requests_house_idx")
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS material_requests (
                id INTEGER PRIMARY KEY,
                house_id TEXT NOT NULL REFERENCES provided_houses(id),
                created_at TEXT NOT NULL,
                created_by TEXT NOT NULL DEFAULT '',
                notes TEXT NOT NULL DEFAULT '',
                data TEXT NOT NULL
            )
            """
        )
        if references_houses:
            db.execute(
                "INSERT INTO material_requests SELECT id, house_id, created_at, created_by, notes, data "
                "FROM material_requests_old"
            )
            db.execute("DROP TABLE material_requests_old")
        db.execute(
            "CREATE INDEX IF NOT EXISTS material_requests_house_idx ON material_requests(house_id, id)"
        )
        if "department" not in movement_columns:
            db.execute(
                "ALTER TABLE movements ADD COLUMN department TEXT NOT NULL DEFAULT ''"
            )
        if "recorded_at" not in movement_columns:
            db.execute(
                "ALTER TABLE movements ADD COLUMN recorded_at TEXT NOT NULL DEFAULT ''"
            )
        # Open houses must hold what was delivered: drop close-outs left behind by earlier reopen actions.
        db.execute(
            "DELETE FROM movements WHERE kind = 'use' AND notes = ? "
            "AND house_id IN (SELECT id FROM houses WHERE is_finished = 0)",
            (CLOSEOUT_NOTE,),
        )
        db.execute("CREATE INDEX IF NOT EXISTS movements_house_idx ON movements(house_id)")
        db.execute(
            "CREATE INDEX IF NOT EXISTS movements_delivery_idx ON movements(house_id, delivery_number)"
        )
        user_columns = {row["name"] for row in db.execute("PRAGMA table_info(app_users)")}
        if "credential_type" not in user_columns:
            db.execute(
                "ALTER TABLE app_users ADD COLUMN credential_type TEXT NOT NULL DEFAULT 'password'"
            )
    _migrate_surveyor_role()


def _migrate_surveyor_role():
    with connect_db() as db:
        schema = db.execute(
            "SELECT sql FROM sqlite_master WHERE type = 'table' AND name = 'app_users'"
        ).fetchone()[0]
        if 'surveyor' in schema:
            return
        db.execute("PRAGMA foreign_keys = OFF")
        db.execute("BEGIN IMMEDIATE")
        try:
            db.execute("""
                CREATE TABLE app_users_new (
                    id INTEGER PRIMARY KEY,
                    username TEXT NOT NULL UNIQUE COLLATE NOCASE,
                    display_name TEXT NOT NULL,
                    role TEXT NOT NULL CHECK (role IN ('admin', 'employee', 'surveyor')),
                    password_salt BLOB NOT NULL,
                    password_hash BLOB NOT NULL,
                    password_iterations INTEGER NOT NULL,
                    credential_type TEXT NOT NULL DEFAULT 'pin',
                    is_active INTEGER NOT NULL DEFAULT 1
                )
            """)
            db.execute("INSERT INTO app_users_new SELECT id, username, display_name, role, password_salt, password_hash, password_iterations, credential_type, is_active FROM app_users")
            db.execute("""CREATE TABLE app_sessions_new (
                token_hash TEXT PRIMARY KEY,
                user_id INTEGER NOT NULL REFERENCES app_users_new(id) ON DELETE CASCADE,
                expires_at INTEGER NOT NULL
            )""")
            db.execute("INSERT INTO app_sessions_new SELECT token_hash, user_id, expires_at FROM app_sessions")
            db.execute("DROP TABLE app_sessions")
            db.execute("DROP TABLE app_users")
            db.execute("ALTER TABLE app_users_new RENAME TO app_users")
            db.execute("ALTER TABLE app_sessions_new RENAME TO app_sessions")
            db.execute("CREATE INDEX app_sessions_expiry_idx ON app_sessions(expires_at)")
            db.commit()
        except Exception:
            db.rollback()
            raise
        finally:
            db.execute("PRAGMA foreign_keys = ON")


def _validate_user_details(username, display_name):
    username = str(username or "").strip()
    display_name = str(display_name or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9._-]{3,64}", username):
        raise ApiError(
            400, "Username must be 3-64 characters: letters, numbers, dot, dash or underscore."
        )
    if not display_name or len(display_name) > 120:
        raise ApiError(400, "Enter a display name (maximum 120 characters).")
    return username, display_name


def _validate_pin(pin):
    if not isinstance(pin, str) or not re.fullmatch(r"[0-9]{6}", pin):
        raise ApiError(400, "PIN must be exactly 6 digits.")
    return pin


def _credential_digest(credential, salt, iterations=CREDENTIAL_ITERATIONS):
    return hashlib.pbkdf2_hmac(
        "sha256", credential.encode("utf-8"), salt, iterations
    )


def _insert_user(db, username, display_name, pin, role):
    username, display_name = _validate_user_details(username, display_name)
    if not isinstance(role, str) or role not in {"admin", "employee", "surveyor"}:
        raise ApiError(400, "Choose a valid account role.")
    pin = _validate_pin(pin)
    salt = secrets.token_bytes(16)
    password_hash = _credential_digest(pin, salt)
    cursor = db.execute(
        """
        INSERT INTO app_users
            (username, display_name, role, password_salt, password_hash, password_iterations, credential_type)
        VALUES (?, ?, ?, ?, ?, ?, 'pin')
        """,
        (
            username,
            display_name,
            role,
            salt,
            password_hash,
            CREDENTIAL_ITERATIONS,
        ),
    )
    return cursor.lastrowid


def _public_user(row):
    return {
        "id": row["id"],
        "username": row["username"],
        "display_name": row["display_name"],
        "role": row["role"],
    }


def _new_session(db, user_id):
    db.execute("DELETE FROM app_sessions WHERE expires_at <= ?", (int(datetime.now().timestamp()),))
    token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(token.encode("ascii")).hexdigest()
    expires_at = int(datetime.now().timestamp()) + SESSION_MAX_AGE
    db.execute(
        "INSERT INTO app_sessions(token_hash, user_id, expires_at) VALUES (?, ?, ?)",
        (token_hash, user_id, expires_at),
    )
    return token, token_hash, expires_at


def _verify_credential(row, credential):
    if not isinstance(credential, str) or len(credential) > 1024:
        credential = ""
    if row is None:
        salt = b"material-tracker-dummy-salt"
        expected = _credential_digest(credential, salt)
        hmac.compare_digest(expected, b"\0" * len(expected))
        return False
    actual = _credential_digest(
        credential, row["password_salt"], row["password_iterations"]
    )
    return hmac.compare_digest(actual, row["password_hash"])


def _user_count(db):
    return db.execute("SELECT COUNT(*) FROM app_users").fetchone()[0]


def _login_attempt_key(username):
    normalized = str(username or "").strip().casefold()[:64]
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _check_login_attempts(db, username_hash, now):
    row = db.execute(
        "SELECT locked_until FROM app_login_attempts WHERE username_hash = ?",
        (username_hash,),
    ).fetchone()
    if row and row["locked_until"] > now:
        raise ApiError(429, "Too many incorrect attempts. Try again in 15 minutes.")


def _record_failed_login(db, username_hash, now):
    row = db.execute(
        """
        SELECT failures, window_started_at FROM app_login_attempts
        WHERE username_hash = ?
        """,
        (username_hash,),
    ).fetchone()
    if row is None or now - row["window_started_at"] >= 900:
        failures = 1
        window_started_at = now
    else:
        failures = row["failures"] + 1
        window_started_at = row["window_started_at"]
    locked_until = now + 900 if failures >= 5 else 0
    db.execute(
        """
        INSERT INTO app_login_attempts
            (username_hash, failures, window_started_at, locked_until)
        VALUES (?, ?, ?, ?)
        ON CONFLICT(username_hash) DO UPDATE SET
            failures = excluded.failures,
            window_started_at = excluded.window_started_at,
            locked_until = excluded.locked_until
        """,
        (username_hash, failures, window_started_at, locked_until),
    )


def _clear_login_attempts(db, username_hash):
    db.execute(
        "DELETE FROM app_login_attempts WHERE username_hash = ?",
        (username_hash,),
    )


def location_key(kind, reference=""):
    if kind == "warehouse":
        return "warehouse"
    if kind == "team":
        return team_location(reference)
    if kind not in LOCATION_TYPES or not reference or ":" in reference:
        raise ApiError(400, "Choose a valid team or house location.")
    return f"{kind}:{reference.strip()}"


def team_location(name, department=""):
    name = str(name or "").strip()
    department = str(department or "").strip()
    if not name:
        raise ApiError(400, "Team name is required.")
    if department:
        return f"team:{quote(department, safe='')}|{quote(name, safe='')}"
    if ":" in name:
        raise ApiError(400, "Team names cannot contain a colon.")
    if "|" in name:
        return f"team:|{quote(name, safe='')}"
    return f"team:{name}"


def employee_location(name):
    return f"employee:{quote(str(name or '').strip(), safe='')}"


def canonical_department(db, name):
    name = str(name or "").strip()
    if not name:
        return ""
    existing = db.execute("SELECT DISTINCT department FROM movements").fetchall()
    for row in existing:
        if row["department"].casefold() == name.casefold():
            return row["department"]
    return name


def canonical_team(db, name, department=""):
    name = str(name or "").strip()
    if not name:
        return ""
    existing = db.execute(
        "SELECT team, department FROM movements WHERE team != ''"
    ).fetchall()
    for row in existing:
        if (
            row["team"].casefold() == name.casefold()
            and row["department"].casefold() == department.casefold()
        ):
            return row["team"]
    return name


def parse_quantity(value, label="Quantity", allow_zero=False):
    try:
        qty = float(value)
    except (TypeError, ValueError):
        raise ApiError(400, f"{label} must be a number.")
    if not math.isfinite(qty) or qty < 0 or (qty == 0 and not allow_zero):
        condition = "zero or greater" if allow_zero else "greater than zero"
        raise ApiError(400, f"{label} must be {condition}.")
    return qty


def parse_legacy_date(value, row_number):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    raise ApiError(400, f"MOVEMENTS row {row_number} has an invalid date.")


def known_locations(db):
    locations = {"warehouse"}
    locations.update(
        f"house:{row['id']}" for row in db.execute("SELECT id FROM houses")
    )
    for row in db.execute(
        "SELECT DISTINCT team, department FROM movements WHERE team != ''"
    ):
        locations.add(team_location(row["team"], row["department"]))
    for row in db.execute(
        "SELECT DISTINCT employee_name FROM houses WHERE TRIM(employee_name) != ''"
    ):
        locations.add(employee_location(row["employee_name"]))
    for row in db.execute("SELECT source, destination FROM movements"):
        locations.add(row["source"])
        locations.add(row["destination"])
    for row in db.execute("SELECT location FROM stocktakes"):
        locations.add(row["location"])
    return locations


def balances_for(db, material_id):
    material = db.execute(
        "SELECT opening_qty FROM materials WHERE id = ?", (material_id,)
    ).fetchone()
    if material is None:
        raise ApiError(404, "Material not found.")

    locations = known_locations(db)
    locations.add("warehouse")
    balances = {}
    movements = db.execute(
        "SELECT id, qty, source, destination FROM movements "
        "WHERE material_id = ? ORDER BY id",
        (material_id,),
    ).fetchall()
    for location in locations:
        snapshot = db.execute(
            "SELECT quantity, movement_id FROM stocktakes "
            "WHERE material_id = ? AND location = ? ORDER BY id DESC LIMIT 1",
            (material_id, location),
        ).fetchone()
        if snapshot:
            balance = snapshot["quantity"]
            after_id = snapshot["movement_id"]
        elif location == "warehouse":
            balance = material["opening_qty"]
            after_id = 0
        else:
            balance = 0.0
            after_id = 0

        if balance is not None:
            for movement in movements:
                if movement["id"] <= after_id:
                    continue
                if movement["source"] == location:
                    balance -= movement["qty"]
                if movement["destination"] == location:
                    balance += movement["qty"]
        balances[location] = balance
    return balances


def _valid_reference(db, location):
    kind, _, reference = location.partition(":")
    if location in {"warehouse", "consumed"}:
        return
    if kind == "house":
        exists = db.execute("SELECT 1 FROM houses WHERE id = ?", (reference,)).fetchone()
    elif kind == "team" and reference.strip():
        exists = True
    else:
        exists = False
    if not exists:
        raise ApiError(400, "A selected house or team location is invalid.")


CLOSEOUT_NOTE = "Automatically recorded from remaining house stock when the house was marked finished."


def _record_house_status(db, house_id, status, changed_by):
    changed_at = datetime.now().isoformat(timespec="seconds")
    changed_by = str(changed_by or "").strip()
    db.execute(
        """
        INSERT INTO house_status_history(house_id, status, changed_by, changed_at)
        VALUES (?, ?, ?, ?)
        """,
        (
            house_id,
            status,
            changed_by,
            changed_at,
        ),
    )
    return {
        "status": status,
        "changed_by": changed_by,
        "changed_at": changed_at,
    }


def _house_audit_by_id(db):
    audit = {}
    for row in db.execute(
        """
        SELECT house_id, status, changed_by, changed_at
        FROM house_status_history
        ORDER BY id
        """
    ):
        if row["status"] not in {"ready", "not_started", "active", "closed"}:
            continue
        entry = {
            "type": "status",
            "status": row["status"],
            "changed_by": row["changed_by"],
            "changed_at": row["changed_at"],
        }
        house_audit = audit.setdefault(
            row["house_id"],
            {"status_history": [], "activity_history": [], "last_movement": None},
        )
        house_audit["status_history"].append(entry)
        house_audit["activity_history"].append(entry)

    for row in db.execute(
        """
        SELECT house_id, created_at, created_by
        FROM material_requests
        ORDER BY id
        """
    ):
        audit.setdefault(
            row["house_id"],
            {"status_history": [], "activity_history": [], "last_movement": None},
        )["activity_history"].append(
            {
                "type": "material_request",
                "created_by": row["created_by"],
                "created_at": row["created_at"],
            }
        )

    for row in db.execute(
        "SELECT house_id, updated_by, updated_at FROM house_surveys"
    ):
        audit.setdefault(
            row["house_id"],
            {"status_history": [], "activity_history": [], "last_movement": None},
        )["activity_history"].append(
            {
                "type": "survey",
                "changed_by": row["updated_by"],
                "changed_at": row["updated_at"],
            }
        )

    latest_movements = {}
    for row in db.execute(
        """
        SELECT m.id, m.house_id, m.kind, m.material_id, p.name AS material,
               p.unit, m.qty, m.delivery_number, m.person, m.team, m.department,
               m.happened_on, m.recorded_at, m.source, m.destination,
               COALESCE(h.employee_name, '') AS house_employee
        FROM movements m
        JOIN materials p ON p.id = m.material_id
        LEFT JOIN houses h ON h.id = m.house_id
        WHERE m.house_id IS NOT NULL
           OR m.source LIKE 'house:%'
           OR m.destination LIKE 'house:%'
        ORDER BY m.id
        """
    ):
        movement = dict(row)
        related_house_ids = {
            movement["house_id"],
            movement["source"][6:] if movement["source"].startswith("house:") else None,
            movement["destination"][6:] if movement["destination"].startswith("house:") else None,
        }
        for house_id in filter(None, related_house_ids):
            house_audit = audit.setdefault(
                house_id,
                {"status_history": [], "activity_history": [], "last_movement": None},
            )
            house_employee = movement["house_employee"]
            if not house_employee:
                assigned = db.execute(
                    "SELECT employee_name FROM houses WHERE id = ?",
                    (house_id,),
                ).fetchone()
                house_employee = assigned["employee_name"] if assigned else ""
            house_audit["activity_history"].append(
                {"type": "movement", **movement, "house_employee": house_employee}
            )
            current = latest_movements.get(house_id)
            if current is None or (movement["recorded_at"], movement["id"]) > (
                current["recorded_at"],
                current["id"],
            ):
                latest_movements[house_id] = movement
    for house_id, movement in latest_movements.items():
        audit[house_id]["last_movement"] = movement
    for house_audit in audit.values():
        house_audit["activity_history"].sort(
            key=lambda entry: (
                entry.get("changed_at")
                or entry.get("created_at")
                or entry.get("recorded_at")
                or "",
                entry.get("id") or 0,
            ),
            reverse=True,
        )
    return audit


def _undo_auto_closeouts(db, house_id):
    """Reopening a house puts back the stock that closing it wrote off automatically."""
    cursor = db.execute(
        "DELETE FROM movements WHERE house_id = ? AND kind = 'use' AND notes = ?",
        (house_id, CLOSEOUT_NOTE),
    )
    return cursor.rowcount


def _require_ready_to_activate(db, house_id):
    """A listed SEAI house needs readiness and a recorded survey before activation."""
    row = db.execute(
        "SELECT p.ready_for_start, h.id AS active_id, hs.house_id AS surveyed_id "
        "FROM provided_houses p "
        "LEFT JOIN houses h ON h.id = p.id "
        "LEFT JOIN house_surveys hs ON hs.house_id = p.id WHERE p.id = ?",
        (str(house_id).strip(),),
    ).fetchone()
    if row and row["active_id"] is None:
        if not row["ready_for_start"]:
            raise ApiError(
                409,
                f"House {str(house_id).strip()} is not marked ready for start. "
                "Mark it ready in Houses & teams before activating it.",
            )
        if not row["surveyed_id"]:
            raise ApiError(
                409,
                "Surveyor review is required before activation. "
                "Record the house survey in the Surveyor room first.",
            )


def _upsert_house(db, house_id, postcode="", owner_name="", employee_name=""):
    if not house_id or not str(house_id).strip():
        raise ApiError(400, "House ID is required for this movement.")
    house_id = str(house_id).strip()
    db.execute(
        """
        INSERT INTO houses(id, postcode, owner_name, employee_name, status_changed_at)
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(id) DO UPDATE SET
            postcode = CASE WHEN excluded.postcode = '' THEN houses.postcode ELSE excluded.postcode END,
            owner_name = CASE WHEN excluded.owner_name = '' THEN houses.owner_name ELSE excluded.owner_name END,
            employee_name = CASE WHEN excluded.employee_name = '' THEN houses.employee_name ELSE excluded.employee_name END
        """,
        (
            house_id,
            str(postcode or "").strip(),
            str(owner_name or "").strip(),
            str(employee_name or "").strip(),
            datetime.now().timestamp(),
        ),
    )
    db.execute(
        "INSERT OR IGNORE INTO provided_houses(id, postcode, owner_name, source) "
        "VALUES (?, ?, ?, 'tracker')",
        (house_id, str(postcode or "").strip(), str(owner_name or "").strip()),
    )
    db.execute(
        "UPDATE provided_houses SET ready_for_start = 0 WHERE id = ?", (house_id,)
    )
    db.execute(
        """
        INSERT OR IGNORE INTO house_checklists
            (house_id, scaffolding_status, esb_unclipped, plumber_status)
        SELECT id,
               CASE plan_scaffolding_required
                   WHEN 1 THEN 'not_done'
                   WHEN 0 THEN 'not_necessary'
                   ELSE plan_scaffolding
               END,
               CASE plan_esb_required
                   WHEN 1 THEN 'no'
                   WHEN 0 THEN 'not_esb'
                   ELSE plan_esb
               END,
               CASE plan_plumber_required
                   WHEN 1 THEN 'not_done'
                   WHEN 0 THEN 'not_necessary'
                   ELSE plan_plumber
               END
        FROM provided_houses
        WHERE id = ? AND (
            (plan_scaffolding_required IS NOT NULL
             AND plan_esb_required IS NOT NULL
             AND plan_plumber_required IS NOT NULL)
            OR (plan_scaffolding IS NOT NULL
                AND plan_esb IS NOT NULL
                AND plan_plumber IS NOT NULL)
        )
        """,
        (house_id,),
    )
    return house_id


def _next_delivery_number(db, house_id):
    if not house_id:
        raise ApiError(400, "House ID is required for a delivery.")
    return db.execute(
        """
        SELECT COALESCE(MAX(delivery_number), 0) + 1
        FROM movements
        WHERE house_id = ? AND kind IN ('delivery', 'pickup')
        """,
        (house_id,),
    ).fetchone()[0]


def _record_movement(
    db,
    record,
    check_balance=True,
    delivery_number=None,
    allow_house_updates=True,
):
    kind = str(record.get("kind", "")).strip()
    if kind not in MOVEMENT_TYPES:
        raise ApiError(400, "Choose a valid movement type.")
    happened_on = str(record.get("date") or date.today().isoformat()).strip()
    try:
        datetime.strptime(happened_on, "%Y-%m-%d")
    except ValueError:
        raise ApiError(400, "Date must use YYYY-MM-DD.")

    material_id = record.get("material_id")
    material = db.execute(
        "SELECT id FROM materials WHERE id = ?", (material_id,)
    ).fetchone()
    if material is None:
        raise ApiError(400, "Choose a valid material.")
    qty = parse_quantity(record.get("qty"))
    department = canonical_department(db, record.get("department"))
    team = canonical_team(db, record.get("team"), department)
    person = str(record.get("person") or "").strip()
    notes = str(record.get("notes") or "").strip()
    if len(notes) > 2000:
        raise ApiError(400, "Notes must be 2,000 characters or fewer.")
    house_id = str(record.get("house_id") or "").strip()
    postcode = str(record.get("postcode") or "").strip()
    owner_name = str(record.get("owner_name") or "").strip()
    employee_name = str(record.get("employee_name") or "").strip()

    def movement_house(reference, postcode="", owner="", assigned_employee=""):
        house_reference = str(reference or "").strip()
        if not house_reference:
            raise ApiError(400, "House ID is required for this movement.")
        closed = db.execute(
            "SELECT is_finished FROM houses WHERE id = ? COLLATE NOCASE",
            (house_reference,),
        ).fetchone()
        if closed and closed["is_finished"]:
            raise ApiError(409, f"House {house_reference} is closed. Reopen it before recording movements.")
        _require_ready_to_activate(db, house_reference)
        if not allow_house_updates:
            exists = db.execute(
                "SELECT 1 FROM houses WHERE id = ?", (house_reference,)
            ).fetchone()
            if not exists:
                raise ApiError(403, "Employees can only record movements for existing houses.")
            return house_reference
        return _upsert_house(
            db, house_reference, postcode, owner, assigned_employee
        )

    if kind in {"pickup", "delivery", "use", "return"} and not house_id:
        raise ApiError(400, "House ID is required for this movement.")
    if house_id:
        house_id = movement_house(house_id, postcode, owner_name, employee_name)

    if kind == "collection":
        if not team or not person:
            raise ApiError(400, "Record the responsible person and team for every movement.")
        source, destination = "warehouse", team_location(team, department)
    elif kind == "pickup":
        if not team or not person:
            raise ApiError(400, "Record the responsible person and team for every movement.")
        source, destination = "warehouse", location_key("house", house_id)
    elif kind == "delivery":
        source_kind = record.get("source_kind", "warehouse")
        if source_kind not in {"warehouse", "team"}:
            raise ApiError(400, "A delivery must come from the warehouse or a team.")
        source_department = canonical_department(
            db, record.get("source_department") or department
        )
        source_team = canonical_team(
            db, record.get("source_team") or team, source_department
        )
        source = (
            "warehouse"
            if source_kind == "warehouse"
            else team_location(source_team, source_department)
        )
        destination = location_key("house", house_id)
        team = team or (source_team if source_kind == "team" else "")
        department = department or (source_department if source_kind == "team" else "")
    elif kind in {"use", "return"}:
        source = location_key("house", house_id)
        destination = "consumed" if kind == "use" else "warehouse"
    else:
        source_kind = str(record.get("source_kind") or "")
        destination_kind = str(record.get("destination_kind") or "")
        if source_kind not in {"team", "house"} or destination_kind not in {"team", "house"}:
            raise ApiError(400, "Transfers must be between a team and/or a house.")
        source_reference = str(record.get("source_reference") or "").strip()
        destination_reference = str(record.get("destination_reference") or "").strip()
        if source_kind == "house":
            movement_house(
                source_reference,
                record.get("source_postcode"),
            )
        if destination_kind == "house":
            movement_house(
                destination_reference,
                record.get("destination_postcode"),
            )
        if source_kind == "team":
            source_department = canonical_department(db, record.get("source_department"))
            source_reference = canonical_team(db, source_reference, source_department)
            source = team_location(source_reference, source_department)
        else:
            source = location_key(source_kind, source_reference)
        if destination_kind == "team":
            destination_department = canonical_department(
                db, record.get("destination_department")
            )
            destination_reference = canonical_team(
                db, destination_reference, destination_department
            )
            destination = team_location(destination_reference, destination_department)
        else:
            destination = location_key(destination_kind, destination_reference)
        if source == destination:
            raise ApiError(400, "Choose different source and destination locations.")
        if not house_id:
            if source_kind == "house":
                house_id = source_reference
            elif destination_kind == "house":
                house_id = destination_reference
        team = team or (
            source_reference if source_kind == "team" else (
                destination_reference if destination_kind == "team" else ""
            )
        )
        department = department or (
            source_department if source_kind == "team" else (
                destination_department if destination_kind == "team" else ""
            )
        )

    if not person or (not team and kind != "use"):
        raise ApiError(400, "Record the responsible person and team for every movement.")
    _valid_reference(db, source)
    _valid_reference(db, destination)
    if check_balance:
        available = balances_for(db, material_id).get(source)
        if available is None:
            raise ApiError(
                409,
                "The source balance is unknown. Record a stocktake before moving material from this location.",
            )
        if available + 1e-9 < qty:
            raise ApiError(
                409,
                f"Insufficient stock at the source location (available: {available:g}).",
            )

    cursor = db.execute(
        """
        INSERT INTO movements
            (happened_on, kind, material_id, qty, source, destination, house_id,
             delivery_number, person, team, department, notes, recorded_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            happened_on, kind, material_id, qty, source, destination, house_id or None,
            (
                delivery_number or _next_delivery_number(db, house_id)
                if kind in {"delivery", "pickup"}
                else None
            ),
            person, team, department, notes,
            datetime.now().isoformat(timespec="seconds"),
        ),
    )
    return cursor.lastrowid


def _record_movement_batch(db, payload, allow_house_updates=True):
    if not isinstance(payload, dict):
        raise ApiError(400, "Movement data must be a JSON object.")
    items = payload.get("items")
    if not isinstance(items, list) or not items:
        raise ApiError(400, "Add at least one material to this movement.")
    if len(items) > 200:
        raise ApiError(400, "A movement can include at most 200 materials.")

    batch_delivery_number = None
    if payload.get("kind") in {"delivery", "pickup"}:
        batch_delivery_number = _next_delivery_number(db, payload.get("house_id"))
    if payload.get("kind") == "pickup" and "is_extra_delivery" in payload:
        is_extra_delivery = payload["is_extra_delivery"]
        if not isinstance(is_extra_delivery, bool):
            raise ApiError(400, "Extra delivery must be selected or cleared.")
        if is_extra_delivery and batch_delivery_number == 1:
            raise ApiError(400, "Record the first warehouse pickup before marking an extra delivery.")
        if not is_extra_delivery and batch_delivery_number > 1:
            raise ApiError(400, "Mark this pickup as an extra delivery to record it as delivery #2 or later.")

    common = {key: value for key, value in payload.items() if key != "items"}
    movement_ids = []
    for item in items:
        if not isinstance(item, dict):
            raise ApiError(400, "Each material entry must be an object.")
        record = {
            **common,
            "material_id": item.get("material_id"),
            "qty": item.get("qty"),
        }
        movement_ids.append(
            _record_movement(
                db,
                record,
                delivery_number=batch_delivery_number,
                allow_house_updates=allow_house_updates,
            )
        )
    return movement_ids, batch_delivery_number


def _movement_json(row):
    return {
        "id": row["id"],
        "date": row["happened_on"],
        "recorded_at": row["recorded_at"],
        "kind": row["kind"],
        "material_id": row["material_id"],
        "material": row["material"],
        "qty": row["qty"],
        "source": row["source"],
        "destination": row["destination"],
        "house_id": row["house_id"],
        "house_owner": row["house_owner"],
        "house_employee": row["house_employee"],
        "house_postcode": row["house_postcode"],
        "delivery_number": row["delivery_number"],
        "person": row["person"],
        "team": row["team"],
        "department": row["department"],
        "notes": row["notes"],
    }


def _movement_rows(db, limit=100, query=""):
    query = str(query or "").strip()
    sql = """
        SELECT m.*, p.name AS material, h.owner_name AS house_owner,
            h.employee_name AS house_employee, h.postcode AS house_postcode
        FROM movements m
        JOIN materials p ON p.id = m.material_id
        LEFT JOIN houses h ON h.id = m.house_id
        WHERE m.kind != 'use'
    """
    if query:
        searchable = (
            "m.happened_on", "m.kind", "p.name", "m.source", "m.destination",
            "m.house_id", "h.owner_name", "h.employee_name", "h.postcode",
            "m.person", "m.team", "m.department", "m.notes",
        )
        sql += " AND (" + " OR ".join(
            "instr(lower(COALESCE({}, '')), lower(?)) > 0".format(column)
            for column in searchable
        ) + ")"
        sql += " ORDER BY m.id DESC"
        return db.execute(sql, (query,) * len(searchable)).fetchall()
    return db.execute(sql + " ORDER BY m.id DESC LIMIT ?", (limit,)).fetchall()


def _house_movement_rows(db, house_id):
    location = f"house:{house_id}"
    return db.execute(
        """
        SELECT m.*, p.name AS material, h.owner_name AS house_owner,
            h.employee_name AS house_employee, h.postcode AS house_postcode
        FROM movements m
        JOIN materials p ON p.id = m.material_id
        LEFT JOIN houses h ON h.id = m.house_id
        WHERE m.house_id = ? OR m.source = ? OR m.destination = ?
        ORDER BY m.happened_on DESC, m.id DESC
        """,
        (house_id, location, location),
    ).fetchall()


def _esb_unclipped_status(value):
    if value is None:
        return None
    if value in (1, True, "1", "yes"):
        return "yes"
    if value in (0, False, "0", "no"):
        return "no"
    return value


def _delivered_to_house(db, house_id):
    """Net quantity of each material sent to a house: deliveries in, minus anything sent back out (not consumed)."""
    location = f"house:{house_id}"
    delivered = {}
    for row in db.execute(
        """
        SELECT material_id,
               SUM(CASE WHEN destination = ? THEN qty ELSE 0 END)
             - SUM(CASE WHEN source = ? AND destination != 'consumed' THEN qty ELSE 0 END) AS qty
        FROM movements WHERE destination = ? OR source = ?
        GROUP BY material_id
        """,
        (location, location, location, location),
    ):
        if row["qty"] and row["qty"] > 1e-9:
            delivered[str(row["material_id"])] = row["qty"]
    return delivered


def _snags_by_house(db, include_created_at=False):
    grouped = {}
    for row in db.execute(
        "SELECT id, house_id, text, done, created_at FROM house_snags ORDER BY done, id"
    ):
        snag = {"id": row["id"], "text": row["text"], "done": bool(row["done"])}
        if include_created_at:
            snag["created_at"] = row["created_at"]
        grouped.setdefault(row["house_id"], []).append(snag)
    return grouped


def _validate_survey_payload(payload):
    if not isinstance(payload, dict):
        raise ApiError(400, "Survey data must be an object.")

    def dimensions(key):
        rows = payload.get(key)
        if not isinstance(rows, list) or len(rows) > 100:
            raise ApiError(400, f"{key.capitalize()} must contain at most 100 rows.")
        result = []
        for row in rows:
            if not isinstance(row, dict):
                raise ApiError(400, f"Invalid {key} measurement.")
            try:
                width, height = float(row.get("w")), float(row.get("h"))
                qty = row.get("qty", 1)
                if isinstance(qty, bool) or int(qty) != qty or not 1 <= int(qty) <= 100:
                    raise ValueError
                qty = int(qty)
            except (TypeError, ValueError, OverflowError):
                raise ApiError(400, f"Enter valid dimensions and quantity for {key}.")
            if not math.isfinite(width) or not math.isfinite(height) or not (0 < width <= 100 and 0 < height <= 100):
                raise ApiError(400, f"{key.capitalize()} dimensions must be between 0 and 100 metres.")
            result.append({"w": width, "h": height, "qty": qty})
        return result

    def lengths(key):
        rows = payload.get(key)
        if not isinstance(rows, list) or len(rows) > 200:
            raise ApiError(400, f"{key.capitalize()} must contain at most 200 lengths.")
        result = []
        for value in rows:
            try:
                length = float(value)
            except (TypeError, ValueError, OverflowError):
                raise ApiError(400, f"Enter a valid length for {key}.")
            if not math.isfinite(length) or not (0 < length <= 100):
                raise ApiError(400, f"{key.capitalize()} must be between 0 and 100 metres.")
            result.append(length)
        return result

    notes = str(payload.get("notes") or "").strip()
    if len(notes) > 2000:
        raise ApiError(400, "Survey notes must be 2,000 characters or fewer.")
    raw_perimeter = payload.get("perimeter_m")
    if raw_perimeter in (None, ""):
        perimeter = None
    else:
        try:
            perimeter = float(raw_perimeter)
        except (TypeError, ValueError, OverflowError):
            raise ApiError(400, "Enter a valid surveyed perimeter.")
        if not math.isfinite(perimeter) or not (0 < perimeter <= 10000):
            raise ApiError(400, "Surveyed perimeter must be between 0 and 10,000 metres.")
    survey = {
        "windows": dimensions("windows"),
        "doors": dimensions("doors"),
        "house_corners": lengths("house_corners"),
        "notes": notes,
    }
    if "perimeter_m" in payload:
        survey["perimeter_m"] = perimeter
    return survey


def _dashboard(db):
    snags = _snags_by_house(db)
    house_audit = _house_audit_by_id(db)
    materials = []
    for row in db.execute("SELECT * FROM materials ORDER BY name COLLATE NOCASE"):
        balances = balances_for(db, row["id"])
        warehouse = balances.get("warehouse")
        negative_locations = [
            location
            for location, quantity in balances.items()
            if location != "consumed" and quantity is not None and quantity < -1e-9
        ]
        materials.append(
            {
                "id": row["id"],
                "code": row["code"],
                "name": row["name"],
                "unit": row["unit"],
                "opening_qty": row["opening_qty"],
                "warehouse_qty": warehouse,
                "balances": balances,
                "negative_locations": negative_locations,
                "needs_stocktake": warehouse is None or bool(negative_locations),
            }
        )
    active_house_count = db.execute(
        "SELECT COUNT(*) FROM houses WHERE is_finished = 0"
    ).fetchone()[0]
    provided_house_count = db.execute(
        "SELECT COUNT(*) FROM provided_houses"
    ).fetchone()[0]
    closed_house_count = db.execute(
        "SELECT COUNT(*) FROM houses WHERE is_finished = 1"
    ).fetchone()[0]
    ready_house_count = db.execute(
        "SELECT COUNT(*) FROM provided_houses p WHERE p.ready_for_start = 1 "
        "AND NOT EXISTS (SELECT 1 FROM houses h WHERE h.id = p.id)"
    ).fetchone()[0]
    not_started_house_count = db.execute(
        "SELECT COUNT(*) FROM provided_houses p WHERE p.ready_for_start = 0 "
        "AND NOT EXISTS (SELECT 1 FROM houses h WHERE h.id = p.id)"
    ).fetchone()[0]
    team_house_counts = {}
    for row in db.execute(
        """
        SELECT DISTINCT m.team, m.department, h.id, h.is_finished, h.status_changed_at
        FROM movements m
        JOIN houses h
          ON h.id = m.house_id
          OR m.source = 'house:' || h.id
          OR m.destination = 'house:' || h.id
        WHERE m.team != ''
        ORDER BY h.status_changed_at DESC, h.id COLLATE NOCASE
        """
    ):
        counts = team_house_counts.setdefault(
            team_location(row["team"], row["department"]),
            {"open_house_ids": [], "closed_house_ids": []},
        )
        house_ids = (
            counts["closed_house_ids"]
            if row["is_finished"]
            else counts["open_house_ids"]
        )
        house_ids.append(row["id"])
    for row in db.execute(
        """
        SELECT id, employee_name, is_finished
        FROM houses
        WHERE TRIM(employee_name) != ''
        """
    ):
        counts = team_house_counts.setdefault(
            employee_location(row["employee_name"]),
            {"open_house_ids": [], "closed_house_ids": []},
        )
        house_ids = (
            counts["closed_house_ids"]
            if row["is_finished"]
            else counts["open_house_ids"]
        )
        house_ids.append(row["id"])
    for counts in team_house_counts.values():
        counts["open_house_ids"].sort(key=str.casefold)
        counts["closed_house_ids"].sort(key=str.casefold)
        counts["open"] = len(counts["open_house_ids"])
        counts["closed"] = len(counts["closed_house_ids"])
    return {
        "materials": materials,
        "demo_mode": (
            db.execute(
                "SELECT value FROM app_settings WHERE key = 'demo_mode'"
            ).fetchone() or {"value": "0"}
        )["value"] == "1",
        "houses": active_house_count,
        "house_summary": {
            "provided": provided_house_count,
            "active": active_house_count,
            "closed": closed_house_count,
            "ready": ready_house_count,
            "not_started": not_started_house_count,
        },
        "movements": [_movement_json(row) for row in _movement_rows(db, 12)],
        "teams": sorted(
            {
                " · ".join(filter(None, (row["department"], row["team"])))
                for row in db.execute(
                    "SELECT DISTINCT department, team FROM movements WHERE team != ''"
                )
            },
            key=str.casefold,
        ),
        "team_house_counts": team_house_counts,
        "house_list": [
            {
                **dict(row),
                "esb_unclipped": _esb_unclipped_status(row["esb_unclipped"]),
                "snags": snags.get(row["id"], []),
                **house_audit.get(
                    row["id"], {"status_history": [], "last_movement": None}
                ),
            }
            for row in db.execute(
                """
                SELECT h.id, COALESCE(NULLIF(p.postcode, ''), h.postcode) AS postcode,
                       COALESCE(NULLIF(p.owner_name, ''), h.owner_name) AS owner_name,
                       h.employee_name, p.address, p.source, p.area_m2, p.perimeter_m,
                       p.latitude, p.longitude, h.is_finished,
                       h.stage, h.status_changed_at, COALESCE(p.notes, '') AS notes,
                       COALESCE(hs.updated_by, '') AS survey_updated_by,
                       COALESCE(hs.updated_at, '') AS survey_updated_at,
                       p.plan_scaffolding_required,
                       COALESCE(p.plan_scaffolding_notes, '') AS plan_scaffolding_notes,
                       p.plan_esb_required,
                       COALESCE(p.plan_esb_notes, '') AS plan_esb_notes,
                       p.plan_plumber_required,
                       COALESCE(p.plan_plumber_notes, '') AS plan_plumber_notes,
                       c.scaffolding_status, c.esb_unclipped, c.plumber_status
                FROM houses h
                LEFT JOIN house_checklists c ON c.house_id = h.id
                LEFT JOIN provided_houses p ON p.id = h.id
                LEFT JOIN house_surveys hs ON hs.house_id = h.id
                ORDER BY h.status_changed_at DESC, h.id COLLATE NOCASE
                """
            )
        ],
        "locations": sorted(known_locations(db)),
    }


def _import_house_locations_csv(db, contents):
    try:
        reader = csv.DictReader(StringIO(contents.decode("utf-8-sig")), strict=True)
    except UnicodeDecodeError:
        raise ApiError(400, "House-location CSV must use UTF-8 encoding.")
    if not reader.fieldnames:
        raise ApiError(400, "House-location CSV needs a header row.")
    try:
        records = list(reader)
    except csv.Error:
        raise ApiError(400, "House-location CSV is malformed.")

    aliases = {
        "id": {"id", "houseid", "house"},
        "postcode": {"postcode", "eircode", "irishpostcode"},
        "address": {"address", "propertyaddress"},
        "owner_name": {"owner", "ownername"},
        "area_m2": {"area", "aream2", "floorarea"},
        "perimeter_m": {"perimeter", "perimeterm"},
        "latitude": {"latitude", "lat"},
        "longitude": {"longitude", "lon", "lng", "long"},
    }
    columns = {}
    for raw_name in reader.fieldnames:
        normalized = re.sub(r"[^a-z0-9]", "", str(raw_name or "").casefold())
        for name, accepted in aliases.items():
            if normalized in accepted:
                if name in columns:
                    raise ApiError(400, f"House-location CSV has duplicate {name} columns.")
                columns[name] = raw_name
                break
    required = {"id", "postcode", "latitude", "longitude"}
    if not required.issubset(columns):
        missing = ", ".join(sorted(required - columns.keys()))
        raise ApiError(400, f"House-location CSV is missing required columns: {missing}.")

    rows = []
    house_ids = set()

    def optional_number(raw, label, row_number):
        if raw is None or not str(raw).strip():
            return None
        try:
            value = float(raw)
        except (TypeError, ValueError):
            raise ApiError(400, f"House-location CSV row {row_number} has invalid {label}.")
        if not math.isfinite(value) or value < 0:
            raise ApiError(400, f"House-location CSV row {row_number} has invalid {label}.")
        return value

    if len(records) > 10000:
        raise ApiError(400, "House-location CSV cannot contain more than 10,000 rows.")
    for row_number, record in enumerate(records, start=2):
        house_id = str(record.get(columns["id"]) or "").strip()
        postcode = str(record.get(columns["postcode"]) or "").strip().upper()
        if not house_id or len(house_id) > 100:
            raise ApiError(400, f"House-location CSV row {row_number} needs a valid house ID.")
        if not postcode:
            raise ApiError(400, f"House-location CSV row {row_number} needs an Eircode.")
        if house_id.casefold() in house_ids:
            raise ApiError(400, f"House-location CSV contains duplicate house ID {house_id}.")
        house_ids.add(house_id.casefold())
        try:
            latitude = float(record.get(columns["latitude"], ""))
            longitude = float(record.get(columns["longitude"], ""))
        except (TypeError, ValueError):
            raise ApiError(400, f"House-location CSV row {row_number} needs valid coordinates.")
        if not math.isfinite(latitude) or not -90 <= latitude <= 90:
            raise ApiError(400, f"House-location CSV row {row_number} has invalid latitude.")
        if not math.isfinite(longitude) or not -180 <= longitude <= 180:
            raise ApiError(400, f"House-location CSV row {row_number} has invalid longitude.")
        rows.append(
            {
                "id": house_id,
                "postcode": postcode,
                "address": str(record.get(columns.get("address", "")) or "").strip(),
                "owner_name": str(record.get(columns.get("owner_name", "")) or "").strip(),
                "area_m2": optional_number(
                    record.get(columns.get("area_m2", "")), "area", row_number
                ),
                "perimeter_m": optional_number(
                    record.get(columns.get("perimeter_m", "")), "perimeter", row_number
                ),
                "latitude": latitude,
                "longitude": longitude,
            }
        )
    if not rows:
        raise ApiError(400, "House-location CSV contains no house records.")

    updated = 0
    for house in rows:
        exists = db.execute(
            "SELECT 1 FROM provided_houses WHERE id = ?", (house["id"],)
        ).fetchone()
        updated += bool(exists)
        db.execute(
            """
            INSERT INTO provided_houses
                (id, postcode, address, owner_name, area_m2, perimeter_m,
                 latitude, longitude, source)
            VALUES
                (:id, :postcode, :address, :owner_name, :area_m2, :perimeter_m,
                 :latitude, :longitude, 'government')
            ON CONFLICT(id) DO UPDATE SET
                postcode = excluded.postcode,
                address = COALESCE(NULLIF(excluded.address, ''), provided_houses.address),
                owner_name = COALESCE(NULLIF(excluded.owner_name, ''), provided_houses.owner_name),
                area_m2 = COALESCE(excluded.area_m2, provided_houses.area_m2),
                perimeter_m = COALESCE(excluded.perimeter_m, provided_houses.perimeter_m),
                latitude = excluded.latitude,
                longitude = excluded.longitude,
                source = 'government'
            """,
            house,
        )
    total = db.execute("SELECT COUNT(*) FROM provided_houses").fetchone()[0]
    return {"imported": len(rows) - updated, "updated": updated, "total": total}


def _import_workbook(db, path):
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        raise ApiError(501, "Excel import requires openpyxl. Install dependencies from requirements.txt.")

    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except (OSError, ValueError, zipfile.BadZipFile, ParseError, InvalidFileException) as exc:
        raise ApiError(400, f"Could not read this Excel workbook: {exc}")
    try:
        return _import_workbook_contents(db, workbook)
    finally:
        workbook.close()


def _import_workbook_contents(db, workbook):
    if "STOCK" not in workbook.sheetnames or "MOVEMENTS" not in workbook.sheetnames:
        raise ApiError(400, "Workbook must contain STOCK and MOVEMENTS sheets.")
    if db.execute("SELECT 1 FROM materials LIMIT 1").fetchone():
        raise ApiError(409, "Import is only allowed into an empty tracker to prevent duplicate records.")

    stock = workbook["STOCK"]
    movement_sheet = workbook["MOVEMENTS"]
    stock_headers = {
        str(stock.cell(1, column).value or "").strip().casefold(): column
        for column in range(1, stock.max_column + 1)
    }
    movement_headers = {
        str(movement_sheet.cell(1, column).value or "").strip().casefold(): column
        for column in range(1, movement_sheet.max_column + 1)
    }
    required_stock = {"material", "starting stock"}
    required_movements = {"date", "house id", "material", "qty", "movement"}
    if not required_stock.issubset(stock_headers) or not required_movements.issubset(
        movement_headers
    ):
        raise ApiError(400, "Workbook columns do not match the expected tracker layout.")

    warnings = []
    material_ids = {}
    used_names = set()
    for row_number in range(2, stock.max_row + 1):
        name = stock.cell(row_number, stock_headers["material"]).value
        if name is None or not str(name).strip():
            continue
        name = str(name).strip()
        normalized = name.casefold()
        if normalized in used_names:
            raise ApiError(400, f"Duplicate material name in STOCK row {row_number}.")
        used_names.add(normalized)
        raw_qty = stock.cell(row_number, stock_headers["starting stock"]).value
        opening_qty = None
        if raw_qty is not None and str(raw_qty).strip() not in {"", "?"}:
            opening_qty = parse_quantity(raw_qty, f"STOCK row {row_number} starting stock", True)
        else:
            warnings.append(f"{name}: starting stock is unknown; record a stocktake before collection.")
        code_col = stock_headers.get("material id")
        code = stock.cell(row_number, code_col).value if code_col else None
        cursor = db.execute(
            "INSERT INTO materials(code, name, opening_qty) VALUES (?, ?, ?)",
            (str(code).strip() if code is not None else "", name, opening_qty),
        )
        material_ids[normalized] = cursor.lastrowid

    def value(row_number, column):
        index = movement_headers.get(column)
        return movement_sheet.cell(row_number, index).value if index else None

    missing_accountability = 0
    skipped_zero_movements = 0
    for row_number in range(2, movement_sheet.max_row + 1):
        name = value(row_number, "material")
        if name is None or not str(name).strip():
            continue
        name = str(name).strip()
        raw_qty = value(row_number, "qty")
        try:
            qty = parse_quantity(
                raw_qty, f"MOVEMENTS row {row_number} quantity", allow_zero=True
            )
        except ApiError as exc:
            raise ApiError(400, str(exc))
        if qty == 0:
            skipped_zero_movements += 1
            continue

        happened_on = parse_legacy_date(value(row_number, "date"), row_number)

        original_kind = str(value(row_number, "movement") or "").strip().casefold()
        kind_map = {
            "delivery to house": "delivery",
            "return to warehouse": "return",
            "used on house": "use",
        }
        kind = kind_map.get(original_kind)
        if kind is None:
            raise ApiError(400, f"MOVEMENTS row {row_number} has an unsupported movement type.")
        house_id = value(row_number, "house id")
        if house_id is None or not str(house_id).strip():
            raise ApiError(400, f"MOVEMENTS row {row_number} is missing its House ID.")
        house_id = str(house_id).strip()
        normalized = name.casefold()
        material_id = material_ids.get(normalized)
        if material_id is None:
            cursor = db.execute(
                "INSERT INTO materials(code, name, opening_qty) VALUES ('', ?, 0)",
                (name,),
            )
            material_id = cursor.lastrowid
            material_ids[normalized] = material_id
            warnings.append(f"{name}: no STOCK row found; imported with zero opening stock.")
        _upsert_house(
            db,
            house_id,
            value(row_number, "postcode / eircode"),
            value(row_number, "house owner name"),
        )
        source = location_key("house", house_id) if kind in {"return", "use"} else "warehouse"
        destination = "warehouse" if kind == "return" else (
            "consumed" if kind == "use" else location_key("house", house_id)
        )
        person = str(value(row_number, "delivered by") or "").strip()
        team = str(value(row_number, "worker team") or "").strip()
        if not person or not team:
            missing_accountability += 1
        notes = str(value(row_number, "notes") or "").strip()
        db.execute(
            """
            INSERT INTO movements
                (happened_on, kind, material_id, qty, source, destination, house_id, delivery_number, person, team, department, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                happened_on, kind, material_id, qty, source, destination, house_id,
                _next_delivery_number(db, house_id) if kind == "delivery" else None,
                person, team, "", notes,
            ),
        )

    if missing_accountability:
        warnings.append(
            f"{missing_accountability} historical movement(s) have no person or team recorded."
        )
    if skipped_zero_movements:
        warnings.append(
            f"{skipped_zero_movements} zero-quantity movement row(s) were skipped."
        )
    negative_balances = 0
    for row in db.execute("SELECT id FROM materials"):
        for location, quantity in balances_for(db, row["id"]).items():
            if location != "consumed" and quantity is not None and quantity < -1e-9:
                negative_balances += 1
    if negative_balances:
        warnings.append(
            f"{negative_balances} warehouse, team or house balance(s) are negative; reconcile them with physical stocktakes."
        )
    return {
        "material_count": len(material_ids),
        "movement_count": db.execute("SELECT COUNT(*) FROM movements").fetchone()[0],
        "house_count": db.execute("SELECT COUNT(*) FROM houses").fetchone()[0],
        "warnings": warnings,
    }


class RequestHandler(BaseHTTPRequestHandler):
    server_version = "EWITracker/1.0"

    def log_message(self, format_string, *args):
        print("%s - %s" % (self.address_string(), format_string % args), file=sys.stderr)

    def _send_json(self, status, payload, headers=()):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        for name, value in headers:
            self.send_header(name, value)
        self.end_headers()
        self.wfile.write(body)

    def _send_file(self, path, content_type):
        try:
            body = path.read_bytes()
        except OSError:
            self._send_json(404, {"error": "Page not found."})
            return
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(body)

    def _send_pdf(self, filename, body, attachment=False):
        self.send_response(200)
        self.send_header("Content-Type", "application/pdf")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header(
            "Content-Disposition",
            f"{'attachment' if attachment else 'inline'}; filename*=UTF-8''{quote(filename, safe='')}",
        )
        self.end_headers()
        self.wfile.write(body)

    def _send_png(self, body, headers=()):
        self.send_response(200)
        self.send_header("Content-Type", "image/png")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        for name, value in headers:
            self.send_header(name, value)
        self.end_headers()
        self.wfile.write(body)

    def _body(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            raise ApiError(400, "Invalid request body.")
        if length <= 0 or length > 10_000_000:
            raise ApiError(400, "Request body is empty or too large.")
        try:
            return self.rfile.read(length)
        except OSError:
            raise ApiError(400, "Could not read request body.")

    def _json_body(self):
        try:
            return json.loads(self._body())
        except (json.JSONDecodeError, UnicodeDecodeError):
            raise ApiError(400, "Request body must be valid JSON.")

    def _session_cookie(self):
        cookie = SimpleCookie()
        try:
            cookie.load(self.headers.get("Cookie", ""))
        except CookieError:
            return ""
        morsel = cookie.get("ewi_session")
        return morsel.value if morsel else ""

    def _current_user(self):
        token = self._session_cookie()
        if not token:
            return None
        token_hash = hashlib.sha256(token.encode("ascii", errors="ignore")).hexdigest()
        now = int(datetime.now().timestamp())
        with connect_db() as db:
            return db.execute(
                """
                SELECT u.id, u.username, u.display_name, u.role, s.token_hash
                FROM app_sessions s
                JOIN app_users u ON u.id = s.user_id
                WHERE s.token_hash = ? AND s.expires_at > ? AND u.is_active = 1
                """,
                (token_hash, now),
            ).fetchone()

    def _login_required(self):
        user = self._current_user()
        if user is None:
            self._send_json(401, {"error": "Sign in to continue."})
            return None
        if user["role"] == "surveyor":
            allowed = self.command == "GET" and self.path.split("?", 1)[0] == "/api/calculator/houses"
            allowed = allowed or (
                self.command == "GET"
                and (
                    urlparse(self.path).path == "/api/house-documents"
                    or urlparse(self.path).path.startswith("/api/house-documents/")
                    or urlparse(self.path).path == "/api/seai-documents"
                    or urlparse(self.path).path.startswith("/api/seai-documents/")
                )
            )
            allowed = allowed or (
                self.command == "POST"
                and urlparse(self.path).path
                in {
                    "/api/house-survey",
                    "/api/house-notes",
                    "/api/survey-notes",
                    "/api/house-work-plan",
                    "/api/seai-measure-variation",
                    "/api/seai-documents/share",
                    "/api/auth/pin",
                }
            )
            allowed = allowed or (
                self.command == "POST"
                and urlparse(self.path).path == "/api/house-documents"
            )
            allowed = allowed or (
                self.command == "POST"
                and re.fullmatch(
                    r"/api/house-documents/\d+/annotations",
                    urlparse(self.path).path,
                )
                is not None
            )
            allowed = allowed or (
                self.command == "POST"
                and re.fullmatch(
                    r"/api/seai-documents/\d+/annotations",
                    urlparse(self.path).path,
                )
                is not None
            )
            allowed = allowed or (
                self.command == "POST"
                and re.fullmatch(
                    r"/api/seai-documents/\d+/share",
                    urlparse(self.path).path,
                )
                is not None
            )
            if not allowed:
                self._send_json(403, {"error": "Surveyor access is limited to ready-house surveys."})
                return None
        return user

    def _admin_required(self, user):
        if user["role"] != "admin":
            self._send_json(403, {"error": "Administrator access is required."})
            return False
        return True

    def _session_headers(self, token, max_age=SESSION_MAX_AGE):
        attributes = [
            f"ewi_session={token}",
            "Path=/",
            "HttpOnly",
            "SameSite=Strict",
            f"Max-Age={max_age}",
        ]
        if COOKIE_SECURE:
            attributes.append("Secure")
        return [("Set-Cookie", "; ".join(attributes))]

    def _expired_session_headers(self):
        return self._session_headers("", 0)

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/":
            self._send_file(STATIC_DIR / "index.html", "text/html; charset=utf-8")
        elif path == "/oldest-house-test":
            user = self._login_required()
            if user is None:
                return
            if user["role"] not in {"admin", "surveyor"}:
                self._send_json(
                    403,
                    {"error": "Surveyor access is required to view house photos."},
                )
                return
            self._send_file(
                STATIC_DIR / "oldest-house-test.html", "text/html; charset=utf-8"
            )
        elif path.startswith("/oldest-house-test/photos/"):
            user = self._login_required()
            if user is None:
                return
            if user["role"] not in {"admin", "surveyor"}:
                self._send_json(
                    403,
                    {"error": "Surveyor access is required to view house photos."},
                )
                return
            photo_name = path.rsplit("/", 1)[-1]
            photos = {
                "front.jpg": STATIC_DIR / "oldest-house-test" / "front.jpg",
                "rear.jpg": STATIC_DIR / "oldest-house-test" / "rear.jpg",
                "gable.jpg": STATIC_DIR / "oldest-house-test" / "gable.jpg",
            }
            photo = photos.get(photo_name)
            if photo is None:
                self._send_json(404, {"error": "Photo not found."})
                return
            self._send_file(photo, "image/jpeg")
        elif path.startswith("/oldest-house-test/pdfs/"):
            user = self._login_required()
            if user is None:
                return
            if user["role"] not in {"admin", "surveyor"}:
                self._send_json(
                    403,
                    {"error": "Surveyor access is required to view house documents."},
                )
                return
            pdf_name = path.rsplit("/", 1)[-1]
            pdfs = {
                "application-letter.pdf": (
                    "Call Off Letter - WH570150.pdf",
                    STATIC_DIR / "oldest-house-test" / "pdfs" / "application-letter.pdf",
                ),
                "survey-report.pdf": (
                    "Surveyor - WH570150 Survey Report_11zon.pdf",
                    STATIC_DIR / "oldest-house-test" / "pdfs" / "survey-report.pdf",
                ),
                "survey-sketch.pdf": (
                    "Surveyor - WH570150 Survey Sketch.pdf",
                    STATIC_DIR / "oldest-house-test" / "pdfs" / "survey-sketch.pdf",
                ),
                "floor-plan-1.pdf": (
                    "Installer - Floor Plan_20260915_151443_308.pdf",
                    STATIC_DIR / "oldest-house-test" / "pdfs" / "floor-plan-1.pdf",
                ),
                "floor-plan-2.pdf": (
                    "Installer - Floor Plan_20260915_151556_411.pdf",
                    STATIC_DIR / "oldest-house-test" / "pdfs" / "floor-plan-2.pdf",
                ),
                "floor-plan-3.pdf": (
                    "Installer - Floor Plan_20260915_151616_123.pdf",
                    STATIC_DIR / "oldest-house-test" / "pdfs" / "floor-plan-3.pdf",
                ),
            }
            pdf = pdfs.get(pdf_name)
            if pdf is None:
                self._send_json(404, {"error": "PDF not found."})
                return
            try:
                content = pdf[1].read_bytes()
            except OSError:
                self._send_json(404, {"error": "PDF not found."})
                return
            self._send_pdf(pdf[0], content)
        elif path == "/app.css":
            self._send_file(STATIC_DIR / "app.css", "text/css; charset=utf-8")
        elif path == "/app.js":
            self._send_file(STATIC_DIR / "app.js", "text/javascript; charset=utf-8")
        elif path == "/brusna-curious-cat-logo.png":
            self._send_file(STATIC_DIR / "brusna-curious-cat-logo.png", "image/png")
        elif path == "/calculator":
            self._send_file(ROOT / "Auto_calc.html", "text/html; charset=utf-8")
        elif path == "/api/calculator/houses":
            user = self._login_required()
            if user is None:
                return
            params = parse_qs(urlparse(self.path).query)
            query = params.get("q", [""])[0].strip()
            status_filter = params.get("status", [""])[0]
            if user["role"] == "surveyor":
                status_filter = "ready"
            wanted = {
                "not_started": "h.id IS NULL AND p.ready_for_start = 0",
                "ready": "h.id IS NULL AND p.ready_for_start = 1",
                "pending": "h.id IS NULL",
                "all": "1 = 1",
                "map": "((h.id IS NOT NULL AND h.is_finished = 0) OR "
                       "(h.id IS NULL AND (p.ready_for_start = 1 OR "
                       "(p.latitude IS NOT NULL AND p.longitude IS NOT NULL))))",
            }.get(status_filter)
            if status_filter == "ready":
                ordering = """
                    (SELECT changed_at FROM house_status_history
                     WHERE house_id = p.id AND status = 'ready'
                     ORDER BY id DESC LIMIT 1) DESC,
                    p.id COLLATE NOCASE
                """
            else:
                ordering = "(h.is_finished = 1), p.id COLLATE NOCASE"
            like = "%" + query.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
            with connect_db() as db:
                rows = db.execute(
                    f"""
                    SELECT p.id, p.postcode, p.address, p.owner_name,
                           COALESCE(h.employee_name, '') AS employee_name,
                           p.area_m2, p.perimeter_m, p.latitude, p.longitude,
                           COALESCE(h.stage, 1) AS stage, p.notes, p.source,
                           p.plan_scaffolding, p.plan_esb, p.plan_plumber,
                           p.plan_scaffolding_required, p.plan_scaffolding_notes,
                           p.plan_esb_required, p.plan_esb_notes,
                           p.plan_plumber_required, p.plan_plumber_notes,
                           hs.data AS survey_data, hs.notes AS survey_notes,
                           hs.updated_by AS survey_updated_by, hs.updated_at AS survey_updated_at,
                           CASE WHEN h.id IS NULL AND p.ready_for_start = 1 THEN 'ready'
                                WHEN h.id IS NULL THEN 'not_started'
                                WHEN h.is_finished = 1 THEN 'closed'
                                ELSE 'active' END AS status
                    FROM provided_houses p
                    LEFT JOIN houses h ON h.id = p.id
                    LEFT JOIN house_surveys hs ON hs.house_id = p.id
                    WHERE (p.id LIKE ? ESCAPE '\\' OR p.postcode LIKE ? ESCAPE '\\'
                       OR p.owner_name LIKE ? ESCAPE '\\' OR p.address LIKE ? ESCAPE '\\')
                      AND ({wanted or "1 = 1"})
                    ORDER BY {ordering}
                    LIMIT ?
                    """,
                    (
                        like,
                        like,
                        like,
                        like,
                        500 if status_filter in {"all", "map"} or wanted else 20,
                    ),
                ).fetchall()
                seai_details = {}
                if user["role"] in {"admin", "surveyor"} and rows:
                    has_seai_applications = db.execute(
                        "SELECT 1 FROM sqlite_master WHERE type = 'table' "
                        "AND name = 'seai_applications'"
                    ).fetchone()
                    if has_seai_applications:
                        placeholders = ",".join("?" for _ in rows)
                        seai_details = {
                            detail["application_id"]: json.loads(
                                detail["detail_data_json"]
                            )
                            for detail in db.execute(
                                "SELECT application_id, detail_data_json "
                                "FROM seai_applications "
                                f"WHERE application_id IN ({placeholders})",
                                [row["id"] for row in rows],
                            )
                        }
                snags = _snags_by_house(db)
                house_audit = _house_audit_by_id(db)
            houses = []
            for row in rows:
                house = dict(row)
                if user["role"] in {"admin", "surveyor"}:
                    house["seai_details"] = seai_details.get(row["id"])
                raw_survey = house.pop("survey_data", None)
                house["survey"] = json.loads(raw_survey) if raw_survey else None
                house["snags"] = snags.get(row["id"], [])
                house.update(
                    house_audit.get(
                        row["id"], {"status_history": [], "last_movement": None}
                    )
                )
                houses.append(house)
            self._send_json(200, {"houses": houses})
        elif path == "/api/house-documents":
            user = self._login_required()
            if user is None:
                return
            if user["role"] not in {"admin", "surveyor"}:
                self._send_json(403, {"error": "Surveyor access is required to view house documents."})
                return
            house_id = parse_qs(urlparse(self.path).query).get("house_id", [""])[0].strip()
            if not house_id:
                self._send_json(400, {"error": "A house ID is required."})
                return
            with connect_db() as db:
                available = db.execute(
                    """
                    SELECT 1 FROM provided_houses p LEFT JOIN houses h ON h.id = p.id
                    WHERE p.id = ? AND h.id IS NULL
                    """,
                    (house_id,),
                ).fetchone()
                if not available:
                    self._send_json(404, {"error": "House was not found."})
                    return
                rows = db.execute(
                    """
                    SELECT id, filename, size_bytes, uploaded_by, uploaded_at
                    FROM house_documents WHERE house_id = ? ORDER BY uploaded_at DESC, id DESC
                    """,
                    (house_id,),
                ).fetchall()
                not_started = _is_not_started_house(db, house_id)
                if not_started:
                    rows = [
                        row for row in rows
                        if not _is_annotated_pdf_filename(row["filename"])
                    ]
            self._send_json(200, {"house_id": house_id, "documents": [dict(row) for row in rows]})
        elif path == "/api/seai-documents":
            user = self._login_required()
            if user is None:
                return
            if user["role"] not in {"admin", "surveyor"}:
                self._send_json(
                    403,
                    {"error": "Administrator or surveyor access is required to view SEAI PDFs."},
                )
                return
            house_id = parse_qs(urlparse(self.path).query).get("house_id", [""])[0].strip()
            if not house_id:
                self._send_json(400, {"error": "A house ID is required."})
                return
            with connect_db() as db:
                document_table = db.execute(
                    "SELECT 1 FROM sqlite_master WHERE type = 'table' "
                    "AND name = 'seai_documents'"
                ).fetchone()
                if not document_table:
                    self._send_json(200, {"house_id": house_id, "documents": []})
                    return
                columns = {
                    row["name"] for row in db.execute("PRAGMA table_info(seai_documents)")
                }
                if "content" not in columns:
                    self._send_json(
                        503,
                        {"error": "SEAI PDF storage has not been initialized."},
                    )
                    return
                not_started = _is_not_started_house(db, house_id)
                rows = db.execute(
                    """
                    SELECT id, category, filename, length(content) AS size_bytes
                    FROM seai_documents
                    WHERE application_id = ? AND lower(filename) LIKE '%.pdf'
                        AND content IS NOT NULL
                    ORDER BY CASE category
                        WHEN 'Application' THEN 0
                        WHEN 'Survey' THEN 1
                        WHEN 'Variation' THEN 2
                        ELSE 3
                    END, filename COLLATE NOCASE
                    """,
                    (house_id,),
                ).fetchall()
                if not_started:
                    rows = [
                        row for row in rows
                        if not _is_annotated_pdf_filename(row["filename"])
                    ]
            self._send_json(
                200,
                {
                    "house_id": house_id,
                    "documents": [
                        {
                            **dict(row),
                            "is_annotated": _is_annotated_pdf_filename(row["filename"]),
                        }
                        for row in rows
                    ],
                },
            )
        elif path.startswith("/api/seai-documents/"):
            user = self._login_required()
            if user is None:
                return
            if user["role"] not in {"admin", "surveyor"}:
                self._send_json(
                    403,
                    {"error": "Administrator or surveyor access is required to view SEAI PDFs."},
                )
                return
            document_path = path[len("/api/seai-documents/"):].split("/")
            is_preview = len(document_path) == 2 and document_path[1] == "preview"
            is_file = len(document_path) == 2 and document_path[1] == "file"
            if (
                len(document_path) != 2
                or not (is_preview or is_file)
                or not document_path[0].isdigit()
            ):
                self._send_json(404, {"error": "SEAI PDF was not found."})
                return
            with connect_db() as db:
                columns = {
                    row["name"] for row in db.execute("PRAGMA table_info(seai_documents)")
                }
                if "content" not in columns:
                    self._send_json(
                        503,
                        {"error": "SEAI PDF storage has not been initialized."},
                    )
                    return
                row = db.execute(
                    """
                    SELECT application_id, filename, content FROM seai_documents
                    WHERE id = ? AND lower(filename) LIKE '%.pdf'
                        AND content IS NOT NULL
                    """,
                    (int(document_path[0]),),
                ).fetchone()
                if (
                    row is not None
                    and _is_not_started_house(db, row["application_id"])
                    and _is_annotated_pdf_filename(row["filename"])
                ):
                    row = None
            if row is None:
                self._send_json(404, {"error": "SEAI PDF was not found."})
                return
            if is_preview:
                page_values = parse_qs(urlparse(self.path).query).get("page", ["1"])
                try:
                    page_number = int(page_values[0])
                except ValueError:
                    self._send_json(400, {"error": "A valid PDF page number is required."})
                    return
                image, page_count = _render_pdf_page(
                    row["content"], page_number
                )
                self._send_png(
                    image, (("X-PDF-Page-Count", str(page_count)),)
                )
            else:
                self._send_pdf(row["filename"], row["content"])
        elif path.startswith("/api/house-documents/"):
            user = self._login_required()
            if user is None:
                return
            if user["role"] not in {"admin", "surveyor"}:
                self._send_json(403, {"error": "Surveyor access is required to view house documents."})
                return
            document_path = path[len("/api/house-documents/"):].split("/")
            is_share_pdf = len(document_path) == 2 and document_path[1] == "share"
            is_preview_page = len(document_path) == 2 and document_path[1] == "preview"
            if len(document_path) not in {1, 2} or (
                len(document_path) == 2 and not (is_share_pdf or is_preview_page)
            ):
                self._send_json(404, {"error": "Document was not found."})
                return
            raw_document_id = document_path[0]
            if not raw_document_id.isdigit():
                self._send_json(404, {"error": "Document was not found."})
                return
            with connect_db() as db:
                row = db.execute(
                    """
                    SELECT d.filename, d.content, d.house_id, hs.data AS survey_data
                    FROM house_documents d
                    JOIN provided_houses p ON p.id = d.house_id
                    LEFT JOIN houses h ON h.id = p.id
                    LEFT JOIN house_surveys hs ON hs.house_id = p.id
                    WHERE d.id = ? AND h.id IS NULL
                    """,
                    (int(raw_document_id),),
                ).fetchone()
                not_started = (
                    row is not None and _is_not_started_house(db, row["house_id"])
                )
                if (
                    not_started
                    and row is not None
                    and _is_annotated_pdf_filename(row["filename"])
                ):
                    row = None
            if row is None:
                self._send_json(404, {"error": "Document was not found."})
                return
            if is_share_pdf and not_started:
                self._send_json(
                    409,
                    {"error": "PDFs are view-only while the house is Not started."},
                )
                return
            if is_preview_page:
                try:
                    page_number = int(parse_qs(urlparse(self.path).query).get("page", ["1"])[0])
                except ValueError:
                    self._send_json(400, {"error": "PDF page number must be a positive integer."})
                    return
                try:
                    page_image, page_count = _render_pdf_page(
                        row["content"], page_number
                    )
                except ApiError as exc:
                    self._send_json(exc.status, {"error": str(exc)})
                    return
                self._send_png(page_image, (("X-PDF-Page-Count", str(page_count)),))
                return
            if not is_share_pdf:
                self._send_pdf(row["filename"], row["content"])
                return
            survey = json.loads(row["survey_data"]) if row["survey_data"] else {}
            survey_note = str(survey.get("notes") or "").strip()
            if not survey_note:
                self._send_json(
                    409,
                    {"error": "Save a survey note before sharing the PDF with notes."},
                )
                return
            try:
                combined_pdf = _create_house_document_share_pdf(
                    row["content"],
                    row["filename"],
                    row["house_id"],
                    survey_note,
                )
            except ApiError as exc:
                self._send_json(exc.status, {"error": str(exc)})
                return
            share_filename = (
                f"{row['house_id']}-{Path(row['filename']).stem}-with-survey-note.pdf"
            )
            self._send_pdf(share_filename, combined_pdf, attachment=True)
        elif path == "/api/calculator/requests":
            user = self._login_required()
            if user is None:
                return
            house_id = parse_qs(urlparse(self.path).query).get("house_id", [""])[0].strip()
            with connect_db() as db:
                rows = db.execute(
                    """
                    SELECT id, house_id, created_at, created_by, notes, data
                    FROM material_requests WHERE house_id = ? ORDER BY id DESC LIMIT 50
                    """,
                    (house_id,),
                ).fetchall()
                delivered = _delivered_to_house(db, house_id)
                next_delivery_number = _next_delivery_number(db, house_id) if house_id else 1
            self._send_json(
                200,
                {
                    "requests": [{**dict(r), "data": json.loads(r["data"])} for r in rows],
                    "delivered": delivered,
                    "next_delivery_number": next_delivery_number,
                },
            )
        elif path == "/brusna-logo.jpg":
            self._send_file(STATIC_DIR / "brusna-logo.jpg", "image/jpeg")
        elif path == "/tommy.jpeg":
            self._send_file(ROOT / "tommy.jpeg", "image/jpeg")
        elif path == "/api/auth/status":
            user = self._current_user()
            with connect_db() as db:
                setup_required = _user_count(db) == 0
            self._send_json(
                200,
                {
                    "setup_required": setup_required,
                    "user": _public_user(user) if user else None,
                },
            )
        elif path == "/api/dashboard":
            user = self._login_required()
            if user is None:
                return
            with connect_db() as db:
                db.execute("BEGIN")
                dashboard = _dashboard(db)
            self._send_json(200, dashboard)
        elif path == "/api/weather":
            user = self._login_required()
            if user is None:
                return
            try:
                forecast = _met_eireann_county_forecast()
            except ApiError as exc:
                self._send_json(exc.status, {"error": str(exc)})
                return
            self._send_json(200, forecast)
        elif path == "/api/movements":
            user = self._login_required()
            if user is None:
                return
            query = parse_qs(urlparse(self.path).query).get("q", [""])[0]
            with connect_db() as db:
                self._send_json(
                    200,
                    {"movements": [_movement_json(row) for row in _movement_rows(db, 500, query)]},
                )
        elif path.startswith("/api/house-checklist/"):
            user = self._login_required()
            if user is None:
                return
            raw_house_id = path[len("/api/house-checklist/"):]
            if not raw_house_id or "/" in raw_house_id:
                self._send_json(404, {"error": "House was not found."})
                return
            house_id = unquote(raw_house_id).strip()
            with connect_db() as db:
                if db.execute(
                    "SELECT 1 FROM houses WHERE id = ?", (house_id,)
                ).fetchone():
                    checklist = db.execute(
                        """
                        SELECT scaffolding_status, esb_unclipped, plumber_status
                        FROM house_checklists WHERE house_id = ?
                        """,
                        (house_id,),
                    ).fetchone()
                else:
                    checklist = db.execute(
                        """
                        SELECT plan_scaffolding AS scaffolding_status,
                               plan_esb AS esb_unclipped,
                               plan_plumber AS plumber_status
                        FROM provided_houses WHERE id = ?
                        """,
                        (house_id,),
                    ).fetchone()
                    if checklist is None:
                        self._send_json(404, {"error": "House was not found."})
                        return
                    if checklist["scaffolding_status"] is None:
                        checklist = None
            self._send_json(
                200,
                {
                    "house_id": house_id,
                    "scaffolding_status": checklist["scaffolding_status"] if checklist else None,
                    "esb_unclipped": (
                        _esb_unclipped_status(checklist["esb_unclipped"])
                        if checklist
                        else None
                    ),
                    "plumber_status": checklist["plumber_status"] if checklist else None,
                },
            )
        elif path.startswith("/api/houses/"):
            user = self._login_required()
            if user is None:
                return
            if not self._admin_required(user):
                return
            raw_house_id = path[len("/api/houses/"):]
            if not raw_house_id or "/" in raw_house_id:
                self._send_json(404, {"error": "House was not found."})
                return
            house_id = unquote(raw_house_id).strip()
            if not house_id:
                self._send_json(404, {"error": "House was not found."})
                return
            with connect_db() as db:
                house = db.execute(
                    """
                    SELECT p.id, COALESCE(NULLIF(p.postcode, ''), h.postcode) AS postcode,
                           COALESCE(NULLIF(p.owner_name, ''), h.owner_name) AS owner_name,
                           COALESCE(h.employee_name, '') AS employee_name,
                           p.address, p.area_m2, p.perimeter_m,
                           COALESCE(hc.scaffolding_status, p.plan_scaffolding) AS scaffolding_status,
                           COALESCE(hc.esb_unclipped, p.plan_esb) AS esb_unclipped,
                           COALESCE(hc.plumber_status, p.plan_plumber) AS plumber_status,
                           COALESCE(h.is_finished, 0) AS is_finished,
                           CASE WHEN h.id IS NULL AND p.ready_for_start = 1 THEN 'ready'
                                WHEN h.id IS NULL THEN 'not_started'
                                WHEN h.is_finished = 1 THEN 'closed'
                                ELSE 'active' END AS status,
                           COALESCE(p.notes, '') AS notes,
                           hs.data AS survey_data,
                           COALESCE(hs.notes, '') AS survey_notes,
                           COALESCE(hs.updated_by, '') AS survey_updated_by,
                           COALESCE(hs.updated_at, '') AS survey_updated_at
                    FROM provided_houses p
                    LEFT JOIN houses h ON h.id = p.id
                    LEFT JOIN house_checklists hc ON hc.house_id = p.id
                    LEFT JOIN house_surveys hs ON hs.house_id = p.id
                    WHERE p.id = ?
                    """,
                    (house_id,),
                ).fetchone()
                if house is None:
                    house = db.execute(
                        """
                        SELECT id, postcode, owner_name, employee_name, NULL AS address,
                               NULL AS area_m2, NULL AS perimeter_m,
                               hc.scaffolding_status, hc.esb_unclipped,
                               hc.plumber_status, is_finished,
                               CASE WHEN is_finished = 1 THEN 'closed' ELSE 'active' END AS status,
                               '' AS notes, NULL AS survey_data, '' AS survey_notes,
                               '' AS survey_updated_by, '' AS survey_updated_at
                        FROM houses
                        LEFT JOIN house_checklists hc ON hc.house_id = houses.id
                        WHERE id = ?
                        """,
                        (house_id,),
                    ).fetchone()
                if house is None:
                    self._send_json(404, {"error": "House was not found."})
                    return
                movements = _house_movement_rows(db, house_id)
                house_data = dict(house)
                try:
                    survey_data = (
                        json.loads(house_data["survey_data"])
                        if house_data.get("survey_data")
                        else {}
                    )
                except (TypeError, json.JSONDecodeError):
                    raise ApiError(500, "Stored house survey data is invalid.")
                if not isinstance(survey_data, dict):
                    raise ApiError(500, "Stored house survey data is invalid.")
                house_data["survey_perimeter_m"] = survey_data.get("perimeter_m")
                house_data.pop("survey_data", None)
                seai_application = db.execute(
                    "SELECT detail_data_json FROM seai_applications "
                    "WHERE application_id = ?",
                    (house_id,),
                ).fetchone() if db.execute(
                    "SELECT 1 FROM sqlite_master WHERE type = 'table' "
                    "AND name = 'seai_applications'"
                ).fetchone() else None
                if seai_application:
                    try:
                        house_data["seai_details"] = json.loads(
                            seai_application["detail_data_json"]
                        )
                    except (TypeError, json.JSONDecodeError):
                        raise ApiError(500, "Stored SEAI application details are invalid.")
                else:
                    house_data["seai_details"] = None
                house_data["esb_unclipped"] = _esb_unclipped_status(
                    house_data["esb_unclipped"]
                )
                house_data["snags"] = _snags_by_house(db, include_created_at=True).get(
                    house_id, []
                )
                house_data.update(
                    _house_audit_by_id(db).get(
                        house_id, {"status_history": [], "last_movement": None}
                    )
                )
                self._send_json(
                    200,
                    {
                        "house": house_data,
                        "movements": [_movement_json(row) for row in movements],
                    },
                )
        elif path == "/api/users":
            user = self._login_required()
            if user is None:
                return
            if not self._admin_required(user):
                return
            with connect_db() as db:
                users = db.execute(
                    """
                    SELECT id, username, display_name, role, is_active
                    FROM app_users ORDER BY username COLLATE NOCASE
                    """
                ).fetchall()
            self._send_json(200, {"users": [dict(row) for row in users]})
        else:
            self._send_json(404, {"error": "Page not found."})

    def do_POST(self):
        path = urlparse(self.path).path
        try:
            if path == "/api/auth/setup":
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "Administrator details must be a JSON object.")
                with connect_db() as db:
                    db.execute("BEGIN IMMEDIATE")
                    if _user_count(db):
                        raise ApiError(409, "The initial administrator has already been set up.")
                    user_id = _insert_user(
                        db,
                        payload.get("username"),
                        payload.get("display_name"),
                        payload.get("pin"),
                        "admin",
                    )
                    token, _, _ = _new_session(db, user_id)
                    user = db.execute(
                        "SELECT id, username, display_name, role FROM app_users WHERE id = ?",
                        (user_id,),
                    ).fetchone()
                self._send_json(
                    201, {"user": _public_user(user)}, self._session_headers(token)
                )
                return
            if path == "/api/auth/login":
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "Sign-in details must be a JSON object.")
                username = str(payload.get("username") or "").strip()
                pin = payload.get("pin")
                username_hash = _login_attempt_key(username)
                now = int(datetime.now().timestamp())
                with connect_db() as db:
                    db.execute("BEGIN IMMEDIATE")
                    _check_login_attempts(db, username_hash, now)
                    row = db.execute(
                        "SELECT * FROM app_users WHERE username = ? COLLATE NOCASE",
                        (username,),
                    ).fetchone()
                    if not _verify_credential(row, pin) or not row["is_active"]:
                        _record_failed_login(db, username_hash, now)
                        attempts = db.execute(
                            "SELECT locked_until FROM app_login_attempts WHERE username_hash = ?",
                            (username_hash,),
                        ).fetchone()
                        locked = attempts["locked_until"] > now
                        token = None
                        user = None
                    else:
                        _clear_login_attempts(db, username_hash)
                        token, _, _ = _new_session(db, row["id"])
                        user = _public_user(row)
                if user is None:
                    if locked:
                        raise ApiError(
                            429,
                            "Too many incorrect attempts. Try again in 15 minutes.",
                        )
                    raise ApiError(401, "Username or PIN is incorrect.")
                self._send_json(
                    200, {"user": user}, self._session_headers(token)
                )
                return
            if path == "/api/auth/logout":
                token = self._session_cookie()
                if token:
                    token_hash = hashlib.sha256(
                        token.encode("ascii", errors="ignore")
                    ).hexdigest()
                    with connect_db() as db:
                        db.execute(
                            "DELETE FROM app_sessions WHERE token_hash = ?",
                            (token_hash,),
                        )
                self._send_json(
                    200, {"ok": True}, self._expired_session_headers()
                )
                return
            if path == "/api/auth/pin":
                user = self._login_required()
                if user is None:
                    return
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "PIN details must be a JSON object.")
                current_pin = payload.get("current_pin")
                new_pin = _validate_pin(payload.get("new_pin"))
                with connect_db() as db:
                    row = db.execute(
                        "SELECT * FROM app_users WHERE id = ?", (user["id"],)
                    ).fetchone()
                    if not _verify_credential(row, current_pin):
                        raise ApiError(401, "Current PIN is incorrect.")
                    salt = secrets.token_bytes(16)
                    db.execute(
                        """
                        UPDATE app_users
                        SET password_salt = ?, password_hash = ?, password_iterations = ?,
                            credential_type = 'pin'
                        WHERE id = ?
                        """,
                        (
                            salt,
                            _credential_digest(new_pin, salt),
                            CREDENTIAL_ITERATIONS,
                            user["id"],
                        ),
                    )
                    db.execute(
                        "DELETE FROM app_sessions WHERE user_id = ?",
                        (user["id"],),
                    )
                    token, _, _ = _new_session(db, user["id"])
                self._send_json(
                    200,
                    {"ok": True},
                    self._session_headers(token),
                )
                return

            user = self._login_required()
            if user is None:
                return
            seai_share_match = re.fullmatch(r"/api/seai-documents/(\d+)/share", path)
            if seai_share_match:
                if user["role"] not in {"admin", "surveyor"}:
                    raise ApiError(
                        403,
                        "Administrator or surveyor access is required to share SEAI PDFs.",
                    )
                document_id = int(seai_share_match.group(1))
                with connect_db() as db:
                    source = db.execute(
                        """
                        SELECT id, application_id, filename, content
                        FROM seai_documents
                        WHERE id = ? AND lower(filename) LIKE '%.pdf'
                            AND content IS NOT NULL
                        """,
                        (document_id,),
                    ).fetchone()
                    if source is None:
                        raise ApiError(404, "SEAI PDF was not found.")
                    _require_pdf_editable(db, source["application_id"])
                    survey_row = db.execute(
                        "SELECT data FROM house_surveys WHERE house_id = ?",
                        (source["application_id"],),
                    ).fetchone()
                    if survey_row is None:
                        raise ApiError(
                            409, "Save a survey note before sharing an SEAI PDF."
                        )
                    try:
                        survey_data = json.loads(survey_row["data"])
                    except (json.JSONDecodeError, TypeError) as exc:
                        raise ApiError(500, "The saved house survey could not be read.") from exc
                    if not isinstance(survey_data, dict):
                        raise ApiError(500, "The saved house survey could not be read.")
                    survey_note = str(survey_data.get("notes") or "").strip()
                    if not survey_note:
                        raise ApiError(
                            409, "Save a survey note before sharing an SEAI PDF."
                        )
                shared_pdf = _create_house_document_share_pdf(
                    source["content"],
                    source["filename"],
                    source["application_id"],
                    survey_note,
                )
                share_filename = (
                    f"{source['application_id']}-"
                    f"{Path(source['filename']).stem}-with-survey-note.pdf"
                )
                self._send_pdf(share_filename, shared_pdf, attachment=True)
                return
            if path == "/api/seai-documents/share":
                if user["role"] not in {"admin", "surveyor"}:
                    raise ApiError(
                        403,
                        "Administrator or surveyor access is required to share SEAI PDFs.",
                    )
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "SEAI PDF share details must be a JSON object.")
                house_id = str(payload.get("house_id") or "").strip()
                selections = payload.get("pages")
                if (
                    not house_id
                    or not isinstance(selections, list)
                    or not selections
                    or len(selections) > HOUSE_DOCUMENT_MAX_PAGES
                ):
                    raise ApiError(
                        400,
                        "Choose between 1 and 300 SEAI PDF pages to share.",
                    )
                normalized_selections = []
                selected_keys = set()
                document_ids = set()
                for selection in selections:
                    if not isinstance(selection, dict):
                        raise ApiError(400, "Selected SEAI PDF pages are invalid.")
                    document_id = selection.get("document_id")
                    page_number = selection.get("page")
                    if (
                        isinstance(document_id, bool)
                        or not isinstance(document_id, int)
                        or document_id < 1
                        or isinstance(page_number, bool)
                        or not isinstance(page_number, int)
                        or page_number < 1
                    ):
                        raise ApiError(
                            400, "Selected SEAI PDF pages must have valid document and page numbers."
                        )
                    page_key = (document_id, page_number)
                    if page_key in selected_keys:
                        raise ApiError(400, "A SEAI PDF page was selected more than once.")
                    selected_keys.add(page_key)
                    document_ids.add(document_id)
                    normalized_selections.append(
                        {"document_id": document_id, "page": page_number}
                    )
                with connect_db() as db:
                    _require_pdf_editable(db, house_id)
                    survey_row = db.execute(
                        "SELECT data FROM house_surveys WHERE house_id = ?",
                        (house_id,),
                    ).fetchone()
                    if survey_row is None:
                        raise ApiError(409, "Save a survey note before sharing selected PDF pages.")
                    try:
                        survey_data = json.loads(survey_row["data"])
                    except (json.JSONDecodeError, TypeError) as exc:
                        raise ApiError(500, "The saved house survey could not be read.") from exc
                    if not isinstance(survey_data, dict):
                        raise ApiError(500, "The saved house survey could not be read.")
                    survey_note = str(survey_data.get("notes") or "").strip()
                    if not survey_note:
                        raise ApiError(
                            409, "Save a survey note before sharing selected PDF pages."
                        )
                    placeholders = ",".join("?" for _ in document_ids)
                    rows = db.execute(
                        f"""
                        SELECT id, category, filename, content FROM seai_documents
                        WHERE application_id = ? AND id IN ({placeholders})
                            AND lower(filename) LIKE '%.pdf' AND content IS NOT NULL
                        """,
                        (house_id, *sorted(document_ids)),
                    ).fetchall()
                    annotated_rows = db.execute(
                        """
                        SELECT id, category, filename, content FROM seai_documents
                        WHERE application_id = ?
                            AND lower(filename) LIKE '%.pdf'
                            AND lower(filename) LIKE '%-annotated%.pdf'
                            AND content IS NOT NULL
                        ORDER BY id DESC
                        """,
                        (house_id,),
                    ).fetchall()
                selected_documents = {row["id"]: row for row in rows}
                if len(selected_documents) != len(document_ids):
                    raise ApiError(
                        404, "One or more selected SEAI PDFs were not found for this house."
                    )
                latest_annotated_documents = {}
                for row in annotated_rows:
                    key = (
                        row["category"],
                        _seai_document_source_stem(row["filename"]).casefold(),
                    )
                    latest_annotated_documents.setdefault(key, row)
                documents = dict(selected_documents)
                resolved_selections = []
                resolved_keys = set()
                for selection in normalized_selections:
                    selected_document = selected_documents[selection["document_id"]]
                    source_stem = _seai_document_source_stem(
                        selected_document["filename"]
                    ).casefold()
                    annotated_document = latest_annotated_documents.get(
                        (selected_document["category"], source_stem)
                    )
                    document = annotated_document or selected_document
                    page_key = (document["id"], selection["page"])
                    if page_key in resolved_keys:
                        raise ApiError(
                            400,
                            "A page was selected more than once after resolving annotated copies.",
                        )
                    resolved_keys.add(page_key)
                    documents[document["id"]] = document
                    resolved_selections.append(
                        {"document_id": document["id"], "page": selection["page"]}
                    )
                shared_pdf = _create_seai_document_share_pdf(
                    resolved_selections,
                    documents,
                    house_id,
                    survey_note,
                )
                share_filename = f"{house_id}-selected-seai-pages-with-survey-note.pdf"
                self._send_pdf(share_filename, shared_pdf, attachment=True)
                return
            seai_annotation_match = re.fullmatch(
                r"/api/seai-documents/(\d+)/annotations", path
            )
            if seai_annotation_match:
                if user["role"] not in {"admin", "surveyor"}:
                    raise ApiError(
                        403, "Administrator or surveyor access is required to annotate SEAI PDFs."
                    )
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "PDF annotations must be a JSON object.")
                document_id = int(seai_annotation_match.group(1))
                with connect_db() as db:
                    columns = {
                        row["name"] for row in db.execute("PRAGMA table_info(seai_documents)")
                    }
                    if "content" not in columns:
                        raise ApiError(503, "SEAI PDF storage has not been initialized.")
                    source = db.execute(
                        """
                        SELECT id, application_id, category, filename, content
                        FROM seai_documents
                        WHERE id = ? AND lower(filename) LIKE '%.pdf'
                            AND content IS NOT NULL
                        """,
                        (document_id,),
                    ).fetchone()
                if source is None:
                    raise ApiError(404, "SEAI PDF was not found.")
                with connect_db() as db:
                    _require_pdf_editable(db, source["application_id"])
                annotated_pdf = _annotate_pdf_document(
                    source["content"], payload.get("annotations")
                )
                if len(annotated_pdf) > SEAI_DOCUMENT_MAX_BYTES:
                    raise ApiError(413, "The annotated SEAI PDF exceeds the 20 MB document limit.")
                with connect_db() as db:
                    with db:
                        db.execute("BEGIN IMMEDIATE")
                        current_source = db.execute(
                            """
                            SELECT 1 FROM seai_documents
                            WHERE id = ? AND content = ?
                            """,
                            (document_id, source["content"]),
                        ).fetchone()
                        if current_source is None:
                            raise ApiError(
                                409,
                                "This SEAI PDF changed while you were annotating it. Reopen it before saving.",
                            )
                        _require_pdf_editable(db, source["application_id"])
                        source_stem = Path(source["filename"]).stem
                        annotated_stem = f"{source_stem}-annotated"
                        annotated_filename = f"{annotated_stem}.pdf"
                        suffix = 2
                        while db.execute(
                            """
                            SELECT 1 FROM seai_documents
                            WHERE application_id = ? AND filename = ? COLLATE NOCASE
                            """,
                            (source["application_id"], annotated_filename),
                        ).fetchone():
                            annotated_filename = f"{annotated_stem}-{suffix}.pdf"
                            suffix += 1
                        cursor = db.execute(
                            """
                            INSERT INTO seai_documents
                                (application_id, category, filename, extension, downloaded, content)
                            VALUES (?, ?, ?, '.pdf', 1, ?)
                            """,
                            (
                                source["application_id"],
                                source["category"],
                                annotated_filename,
                                sqlite3.Binary(annotated_pdf),
                            ),
                        )
                self._send_json(
                    201,
                    {
                        "id": cursor.lastrowid,
                        "source_id": document_id,
                        "filename": annotated_filename,
                        "size_bytes": len(annotated_pdf),
                    },
                )
                return
            annotation_match = re.fullmatch(
                r"/api/house-documents/(\d+)/annotations", path
            )
            if annotation_match:
                if user["role"] not in {"admin", "surveyor"}:
                    raise ApiError(403, "Surveyor access is required to annotate house documents.")
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "PDF annotations must be a JSON object.")
                document_id = int(annotation_match.group(1))
                with connect_db() as db:
                    source = db.execute(
                        """
                        SELECT d.house_id, d.filename, d.content
                        FROM house_documents d
                        JOIN provided_houses p ON p.id = d.house_id
                        LEFT JOIN houses h ON h.id = p.id
                        WHERE d.id = ? AND p.ready_for_start = 1 AND h.id IS NULL
                        """,
                        (document_id,),
                    ).fetchone()
                if source is None:
                    raise ApiError(404, "Document was not found.")
                annotated_pdf = _annotate_pdf_document(
                    source["content"], payload.get("annotations")
                )
                if len(annotated_pdf) > HOUSE_DOCUMENT_MAX_BYTES:
                    raise ApiError(413, "The annotated PDF exceeds the 8 MB document limit.")
                annotated_filename = (
                    f"{Path(source['filename']).stem}-annotated.pdf"
                )
                uploaded_at = datetime.now().isoformat(timespec="seconds")
                with connect_db() as db:
                    with db:
                        ready = db.execute(
                            """
                            SELECT 1 FROM provided_houses p LEFT JOIN houses h ON h.id = p.id
                            WHERE p.id = ? AND p.ready_for_start = 1 AND h.id IS NULL
                            """,
                            (source["house_id"],),
                        ).fetchone()
                        if not ready:
                            raise ApiError(409, "PDFs can only be annotated for houses ready for start.")
                        document_count = db.execute(
                            "SELECT COUNT(*) FROM house_documents WHERE house_id = ?",
                            (source["house_id"],),
                        ).fetchone()[0]
                        if document_count >= HOUSE_DOCUMENT_MAX_COUNT:
                            raise ApiError(409, "This house already has the maximum of 25 PDF documents.")
                        cursor = db.execute(
                            """
                            INSERT INTO house_documents
                                (house_id, filename, content, size_bytes, uploaded_by, uploaded_at)
                            VALUES (?, ?, ?, ?, ?, ?)
                            """,
                            (
                                source["house_id"],
                                annotated_filename,
                                annotated_pdf,
                                len(annotated_pdf),
                                f"Annotated by {user['display_name']}",
                                uploaded_at,
                            ),
                        )
                self._send_json(
                    201,
                    {
                        "id": cursor.lastrowid,
                        "house_id": source["house_id"],
                        "filename": annotated_filename,
                        "size_bytes": len(annotated_pdf),
                        "uploaded_by": f"Annotated by {user['display_name']}",
                        "uploaded_at": uploaded_at,
                    },
                )
                return
            if path == "/api/house-survey":
                if user["role"] not in {"admin", "surveyor"}:
                    raise ApiError(403, "Surveyor access is required to record measurements.")
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "Survey data must be an object.")
                house_id = str(payload.get("house_id") or "").strip()
                survey = _validate_survey_payload(payload.get("survey"))
                with connect_db() as db:
                    with db:
                        ready = db.execute("""
                            SELECT 1 FROM provided_houses p LEFT JOIN houses h ON h.id = p.id
                            WHERE p.id = ? AND p.ready_for_start = 1 AND h.id IS NULL
                        """, (house_id,)).fetchone()
                        if not ready:
                            raise ApiError(409, "Survey measurements can only be saved for houses ready for start.")
                        if survey.get("perimeter_m") is None:
                            raise ApiError(
                                400,
                                "Enter a surveyed perimeter greater than 0 and no more than 10,000 metres.",
                            )
                        db.execute("""
                            INSERT INTO house_surveys(house_id, data, notes, updated_by, updated_at)
                            VALUES (?, ?, ?, ?, ?)
                            ON CONFLICT(house_id) DO UPDATE SET data = excluded.data,
                                notes = excluded.notes, updated_by = excluded.updated_by,
                                updated_at = excluded.updated_at
                        """, (house_id, json.dumps(survey, ensure_ascii=False), survey["notes"],
                                  user["display_name"], datetime.now().isoformat(timespec="seconds")))
                self._send_json(200, {"house_id": house_id, "saved": True, "survey": survey})
                return
            if path == "/api/house-notes":
                if user["role"] not in {"admin", "surveyor"}:
                    raise ApiError(403, "Surveyor access is required to edit house notes.")
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "House notes must be a JSON object.")
                house_id = str(payload.get("house_id") or "").strip()
                notes = str(payload.get("notes") or "").strip()
                if not house_id:
                    raise ApiError(400, "A house ID is required.")
                if len(notes) > 2000:
                    raise ApiError(400, "Notes must be 2,000 characters or fewer.")
                with connect_db() as db:
                    with db:
                        ready = db.execute(
                            """
                            SELECT 1 FROM provided_houses p LEFT JOIN houses h ON h.id = p.id
                            WHERE p.id = ? AND p.ready_for_start = 1 AND h.id IS NULL
                            """,
                            (house_id,),
                        ).fetchone()
                        if not ready:
                            raise ApiError(409, "House notes can only be edited for houses ready for start.")
                        db.execute(
                            "UPDATE provided_houses SET notes = ? WHERE id = ?",
                            (notes, house_id),
                        )
                self._send_json(200, {"house_id": house_id, "notes": notes, "saved": True})
                return
            if path == "/api/survey-notes":
                if user["role"] not in {"admin", "surveyor"}:
                    raise ApiError(403, "Surveyor access is required to edit survey notes.")
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "Survey note data must be a JSON object.")
                house_id = str(payload.get("house_id") or "").strip()
                note = str(payload.get("notes") or "").strip()
                if not house_id:
                    raise ApiError(400, "A house ID is required.")
                if len(note) > 2000:
                    raise ApiError(400, "Survey notes must be 2,000 characters or fewer.")
                now = datetime.now().isoformat(timespec="seconds")
                with connect_db() as db:
                    with db:
                        ready = db.execute(
                            """
                            SELECT 1 FROM provided_houses p LEFT JOIN houses h ON h.id = p.id
                            WHERE p.id = ? AND p.ready_for_start = 1 AND h.id IS NULL
                            """,
                            (house_id,),
                        ).fetchone()
                        if not ready:
                            raise ApiError(409, "Survey notes can only be edited for houses ready for start.")
                        existing = db.execute(
                            "SELECT data FROM house_surveys WHERE house_id = ?",
                            (house_id,),
                        ).fetchone()
                        survey = json.loads(existing["data"]) if existing else {
                            "windows": [],
                            "doors": [],
                            "house_corners": [],
                        }
                        survey["notes"] = note
                        db.execute(
                            """
                            INSERT INTO house_surveys(house_id, data, notes, updated_by, updated_at)
                            VALUES (?, ?, ?, ?, ?)
                            ON CONFLICT(house_id) DO UPDATE SET
                                data = excluded.data,
                                notes = excluded.notes,
                                updated_by = excluded.updated_by,
                                updated_at = excluded.updated_at
                            """,
                            (
                                house_id,
                                json.dumps(survey, ensure_ascii=False),
                                note,
                                user["display_name"],
                                now,
                            ),
                        )
                self._send_json(
                    200,
                    {
                        "house_id": house_id,
                        "notes": note,
                        "survey": survey,
                        "updated_by": user["display_name"],
                        "updated_at": now,
                        "saved": True,
                    },
                )
                return
            if path == "/api/house-documents":
                if user["role"] not in {"admin", "surveyor"}:
                    raise ApiError(403, "Surveyor access is required to upload house documents.")
                content_type = self.headers.get("Content-Type", "")
                if not content_type.startswith("multipart/form-data;"):
                    raise ApiError(400, "Upload a PDF using the file picker.")
                body = self._body()
                match = re.search(r'boundary="?([^";]+)"?', content_type)
                if not match:
                    raise ApiError(400, "Invalid PDF upload.")
                try:
                    boundary = match.group(1).encode("ascii")
                except UnicodeEncodeError as exc:
                    raise ApiError(400, "Invalid PDF upload boundary.") from exc
                house_id = ""
                filename = ""
                file_data = None
                for part in body.split(b"--" + boundary):
                    if b"\r\n\r\n" not in part:
                        continue
                    headers, content = part.split(b"\r\n\r\n", 1)
                    disposition = re.search(
                        rb'Content-Disposition:\s*form-data;\s*name="([^"]+)"(?:;\s*filename="([^"]*)")?',
                        headers,
                        re.IGNORECASE,
                    )
                    if not disposition:
                        continue
                    raw_content = content.rsplit(b"\r\n", 1)[0]
                    if disposition.group(1) == b"house_id":
                        try:
                            house_id = raw_content.decode("utf-8").strip()
                        except UnicodeDecodeError as exc:
                            raise ApiError(400, "Invalid house ID.") from exc
                    elif disposition.group(1) == b"file":
                        raw_filename = disposition.group(2) or b""
                        try:
                            filename = Path(raw_filename.decode("utf-8")).name.strip()
                        except (UnicodeDecodeError, ValueError) as exc:
                            raise ApiError(400, "Invalid PDF filename.") from exc
                        file_data = raw_content
                if not house_id:
                    raise ApiError(400, "A house ID is required.")
                if not filename or not filename.lower().endswith(".pdf"):
                    raise ApiError(400, "House documents must be PDF files.")
                if not file_data:
                    raise ApiError(400, "Choose a PDF file to upload.")
                if len(file_data) > HOUSE_DOCUMENT_MAX_BYTES:
                    raise ApiError(400, "PDF files must be 8 MB or smaller.")
                if not file_data.startswith(b"%PDF-"):
                    raise ApiError(400, "The selected file is not a valid PDF.")
                with connect_db() as db:
                    with db:
                        ready = db.execute(
                            """
                            SELECT 1 FROM provided_houses p LEFT JOIN houses h ON h.id = p.id
                            WHERE p.id = ? AND p.ready_for_start = 1 AND h.id IS NULL
                            """,
                            (house_id,),
                        ).fetchone()
                        if not ready:
                            raise ApiError(404, "Ready house was not found.")
                        document_count = db.execute(
                            "SELECT COUNT(*) FROM house_documents WHERE house_id = ?",
                            (house_id,),
                        ).fetchone()[0]
                        if document_count >= HOUSE_DOCUMENT_MAX_COUNT:
                            raise ApiError(409, "This house already has the maximum of 25 PDF documents.")
                        cursor = db.execute(
                            """
                            INSERT INTO house_documents
                                (house_id, filename, content, size_bytes, uploaded_by, uploaded_at)
                            VALUES (?, ?, ?, ?, ?, ?)
                            """,
                            (
                                house_id,
                                filename,
                                file_data,
                                len(file_data),
                                user["display_name"],
                                datetime.now().isoformat(timespec="seconds"),
                            ),
                        )
                self._send_json(
                    201,
                    {
                        "id": cursor.lastrowid,
                        "house_id": house_id,
                        "filename": filename,
                        "size_bytes": len(file_data),
                    },
                )
                return
            if path not in {
                "/api/movements",
                "/api/movements/batch",
                "/api/house-checklist",
                "/api/house-work-plan",
            }:
                if not self._admin_required(user):
                    return
            if path == "/api/users":
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "Employee details must be a JSON object.")
                with connect_db() as db:
                    with db:
                        user_id = _insert_user(
                            db,
                            payload.get("username"),
                            payload.get("display_name"),
                            payload.get("pin"),
                            payload.get("role", "employee"),
                        )
                        created = db.execute(
                            "SELECT id, username, display_name, role FROM app_users WHERE id = ?",
                            (user_id,),
                        ).fetchone()
                self._send_json(201, {"user": _public_user(created)})
            elif path in {"/api/movements", "/api/movements/batch"}:
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "Movement data must be a JSON object.")
                allow_house_updates = user["role"] == "admin"
                if path.endswith("/batch"):
                    with connect_db() as db:
                        with db:
                            db.execute("BEGIN IMMEDIATE")
                            movement_ids, delivery_number = _record_movement_batch(
                                db, payload, allow_house_updates
                            )
                else:
                    if not isinstance(payload, dict):
                        raise ApiError(400, "Movement data must be a JSON object.")
                    with connect_db() as db:
                        with db:
                            db.execute("BEGIN IMMEDIATE")
                            movement_ids = [
                                _record_movement(
                                    db,
                                    payload,
                                    allow_house_updates=allow_house_updates,
                                )
                            ]
                            delivery_number = None
                            if payload.get("kind") == "delivery":
                                delivery_number = db.execute(
                                    "SELECT delivery_number FROM movements WHERE id = ?",
                                    (movement_ids[0],),
                                ).fetchone()[0]
                if path.endswith("/batch"):
                    self._send_json(
                        201,
                        {"ids": movement_ids, "delivery_number": delivery_number},
                    )
                else:
                    self._send_json(
                        201,
                        {"id": movement_ids[0], "delivery_number": delivery_number},
                    )
            elif path == "/api/house-checklist":
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "House checklist must be a JSON object.")
                if user["role"] not in {"admin", "employee"}:
                    raise ApiError(403, "Administrator or employee access is required.")
                house_id = str(payload.get("house_id") or "").strip()
                if not house_id:
                    raise ApiError(400, "House ID is required.")
                status_options = {
                    "scaffolding_status": {"done", "not_done", "not_necessary"},
                    "esb_unclipped": {"yes", "no", "not_esb"},
                    "plumber_status": {"done", "not_done", "not_necessary"},
                }
                supplied_statuses = {
                    name: payload[name] for name in status_options if name in payload
                }
                if not supplied_statuses:
                    raise ApiError(400, "Choose at least one house checklist status.")
                for name, value in supplied_statuses.items():
                    if (
                        not isinstance(value, str)
                        or value not in status_options[name]
                    ):
                        raise ApiError(
                            400, f"Choose a valid {name.replace('_', ' ')}."
                        )
                with connect_db() as db:
                    with db:
                        is_active = db.execute(
                            "SELECT 1 FROM houses WHERE id = ?", (house_id,)
                        ).fetchone()
                        if not is_active:
                            if user["role"] != "admin":
                                raise ApiError(
                                    403,
                                    "Administrator access is required to edit a pre-start checklist.",
                                )
                            if set(supplied_statuses) != set(status_options):
                                raise ApiError(
                                    400, "Choose all pre-start checklist statuses."
                                )
                            planned = db.execute(
                                "UPDATE provided_houses SET plan_scaffolding = ?, "
                                "plan_esb = ?, plan_plumber = ? WHERE id = ?",
                                (
                                    supplied_statuses["scaffolding_status"],
                                    supplied_statuses["esb_unclipped"],
                                    supplied_statuses["plumber_status"],
                                    house_id,
                                ),
                            )
                            if planned.rowcount == 0:
                                raise ApiError(404, "House was not found.")
                            self._send_json(200, {"house_id": house_id, "saved": True})
                            return
                        db.execute(
                            "INSERT OR IGNORE INTO house_checklists (house_id) VALUES (?)",
                            (house_id,),
                        )
                        for name, value in supplied_statuses.items():
                            db.execute(
                                f"UPDATE house_checklists SET {name} = ? WHERE house_id = ?",
                                (value, house_id),
                            )
                self._send_json(200, {"house_id": house_id, "saved": True})
            elif path == "/api/house-work-plan":
                if user["role"] not in {"admin", "surveyor"}:
                    raise ApiError(
                        403, "Surveyor access is required to save a work plan."
                    )
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "House work plan must be a JSON object.")
                house_id = str(payload.get("house_id") or "").strip()
                if not house_id:
                    raise ApiError(400, "House ID is required.")
                required_values = {}
                notes_values = {}
                for trade in ("scaffolding", "esb", "plumber"):
                    required = payload.get(f"{trade}_required")
                    if required is not None and not isinstance(required, bool):
                        raise ApiError(
                            400, f"Choose whether {trade} work is required."
                        )
                    notes = payload.get(f"{trade}_notes", "")
                    if not isinstance(notes, str) or len(notes) > 1000:
                        raise ApiError(
                            400,
                            f"{trade.title()} work notes must be 1,000 characters or fewer.",
                        )
                    required_values[trade] = (
                        None if required is None else int(required)
                    )
                    notes_values[trade] = notes.strip()
                with connect_db() as db:
                    with db:
                        saved = db.execute(
                            """
                            UPDATE provided_houses
                            SET plan_scaffolding_required = ?,
                                plan_scaffolding_notes = ?,
                                plan_esb_required = ?,
                                plan_esb_notes = ?,
                                plan_plumber_required = ?,
                                plan_plumber_notes = ?
                            WHERE id = ? AND ready_for_start = 1
                              AND NOT EXISTS (
                                  SELECT 1 FROM houses
                                  WHERE houses.id = provided_houses.id
                              )
                            """,
                            (
                                required_values["scaffolding"],
                                notes_values["scaffolding"],
                                required_values["esb"],
                                notes_values["esb"],
                                required_values["plumber"],
                                notes_values["plumber"],
                                house_id,
                            ),
                        )
                        if saved.rowcount == 0:
                            exists = db.execute(
                                "SELECT 1 FROM provided_houses WHERE id = ?",
                                (house_id,),
                            ).fetchone()
                            if not exists:
                                raise ApiError(404, "House was not found.")
                            raise ApiError(
                                409,
                                "Work plans can only be edited for Ready for start houses.",
                            )
                self._send_json(200, {"house_id": house_id, "saved": True})
            elif path == "/api/seai-measure-variation":
                if user["role"] not in {"admin", "surveyor"}:
                    raise ApiError(
                        403, "Administrator or surveyor access is required to edit SEAI quantities."
                    )
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "SEAI measure update must be a JSON object.")
                house_id = str(payload.get("house_id") or "").strip()
                measure_index = payload.get("measure_index")
                variation = payload.get("variation")
                if not house_id:
                    raise ApiError(400, "House ID is required.")
                if (
                    not isinstance(measure_index, int)
                    or isinstance(measure_index, bool)
                    or measure_index < 0
                ):
                    raise ApiError(400, "Choose a valid SEAI measure.")
                if not isinstance(variation, str):
                    raise ApiError(400, "Enter a valid variation quantity.")
                variation = variation.strip()
                if len(variation) > 80 or not re.fullmatch(
                    r"(?:\d+(?:\.\d{1,4})?)?", variation
                ):
                    raise ApiError(
                        400, "Variation quantity must be a non-negative number with up to 4 decimal places."
                    )
                if variation and float(variation) > 1_000_000_000:
                    raise ApiError(400, "Variation quantity is too large.")
                with connect_db() as db:
                    with db:
                        db.execute("BEGIN IMMEDIATE")
                        house = db.execute(
                            """
                            SELECT p.ready_for_start,
                                   CASE WHEN h.id IS NULL THEN 0 ELSE 1 END AS is_active
                            FROM provided_houses p
                            LEFT JOIN houses h ON h.id = p.id
                            WHERE p.id = ?
                            """,
                            (house_id,),
                        ).fetchone()
                        if house is None:
                            raise ApiError(404, "House was not found.")
                        if not house["ready_for_start"]:
                            raise ApiError(
                                409,
                                "Survey variations can only be edited for houses ready for start.",
                            )
                        if user["role"] == "surveyor" and (
                            house["is_active"]
                        ):
                            raise ApiError(
                                409,
                                "Surveyors can only edit quantities for Ready for start houses.",
                            )
                        if not db.execute(
                            "SELECT 1 FROM sqlite_master WHERE type = 'table' "
                            "AND name = 'seai_applications'"
                        ).fetchone():
                            raise ApiError(404, "SEAI application details were not found.")
                        application = db.execute(
                            "SELECT detail_data_json FROM seai_applications "
                            "WHERE application_id = ?",
                            (house_id,),
                        ).fetchone()
                        if application is None:
                            raise ApiError(404, "SEAI application details were not found.")
                        try:
                            details = json.loads(application["detail_data_json"])
                        except (TypeError, json.JSONDecodeError):
                            raise ApiError(500, "Stored SEAI application details are invalid.")
                        measures = details.get("Measures") if isinstance(details, dict) else None
                        if (
                            not isinstance(measures, list)
                            or measure_index >= len(measures)
                            or not isinstance(measures[measure_index], dict)
                        ):
                            raise ApiError(404, "SEAI measure was not found.")
                        measures[measure_index]["Survey Qty including Variation"] = variation
                        db.execute(
                            "UPDATE seai_applications SET detail_data_json = ? "
                            "WHERE application_id = ?",
                            (json.dumps(details, ensure_ascii=False), house_id),
                        )
                self._send_json(
                    200,
                    {
                        "house_id": house_id,
                        "measure_index": measure_index,
                        "variation": variation,
                        "seai_details": details,
                    },
                )
            elif path == "/api/houses/snags":
                user = self._login_required()
                if user is None:
                    return
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "Snag details must be a JSON object.")
                with connect_db() as db:
                    with db:
                        if "snag_id" in payload:
                            snag = db.execute(
                                "SELECT id, house_id FROM house_snags WHERE id = ?",
                                (payload.get("snag_id"),),
                            ).fetchone()
                            if snag is None:
                                raise ApiError(404, "Snag was not found.")
                            if payload.get("delete") is True:
                                db.execute("DELETE FROM house_snags WHERE id = ?", (snag["id"],))
                            elif isinstance(payload.get("done"), bool):
                                db.execute(
                                    "UPDATE house_snags SET done = ? WHERE id = ?",
                                    (int(payload["done"]), snag["id"]),
                                )
                            else:
                                raise ApiError(400, "Choose whether the snag is done.")
                            house_id = snag["house_id"]
                        else:
                            house_id = str(payload.get("id") or "").strip()
                            text = str(payload.get("text") or "").strip()
                            if not text:
                                raise ApiError(400, "Describe the snag first.")
                            if len(text) > 500:
                                raise ApiError(400, "A snag must be 500 characters or fewer.")
                            if not db.execute(
                                "SELECT 1 FROM provided_houses WHERE id = ?", (house_id,)
                            ).fetchone():
                                raise ApiError(404, "House was not found.")
                            db.execute(
                                "INSERT INTO house_snags(house_id, text, created_at) VALUES (?, ?, ?)",
                                (house_id, text, datetime.now().timestamp()),
                            )
                    snags = _snags_by_house(db).get(house_id, [])
                self._send_json(200, {"id": house_id, "snags": snags})
            elif path == "/api/houses/stage":
                user = self._login_required()
                if user is None:
                    return
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "Stage details must be a JSON object.")
                house_id = str(payload.get("id") or "").strip()
                stage = payload.get("stage")
                if isinstance(stage, bool) or stage not in (1, 2):
                    raise ApiError(400, "Stage must be 1 or 2.")
                with connect_db() as db:
                    with db:
                        house = db.execute(
                            "SELECT stage FROM houses WHERE id = ?", (house_id,)
                        ).fetchone()
                        if house is None:
                            raise ApiError(404, "House was not found.")
                        if house["stage"] != stage:
                            db.execute(
                                "UPDATE houses SET stage = ? WHERE id = ?",
                                (stage, house_id),
                            )
                self._send_json(
                    200, {"id": house_id, "stage": stage}
                )
            elif path == "/api/houses/plan":
                user = self._login_required()
                if user is None:
                    return
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "House details must be a JSON object.")
                house_id = str(payload.get("id") or "").strip()
                with connect_db() as db:
                    with db:
                        if not db.execute(
                            "SELECT 1 FROM provided_houses WHERE id = ?", (house_id,)
                        ).fetchone():
                            raise ApiError(404, "House was not found.")
                        if "notes" in payload:
                            notes = str(payload.get("notes") or "").strip()
                            if len(notes) > 2000:
                                raise ApiError(400, "Notes must be 2,000 characters or fewer.")
                            db.execute(
                                "UPDATE provided_houses SET notes = ? WHERE id = ?",
                                (notes, house_id),
                            )
                        if "ready" in payload:
                            if not isinstance(payload["ready"], bool):
                                raise ApiError(400, "Choose whether the house is ready for start.")
                            if db.execute(
                                "SELECT 1 FROM houses WHERE id = ?", (house_id,)
                            ).fetchone():
                                raise ApiError(409, f"House {house_id} is already activated.")
                            current = db.execute(
                                "SELECT ready_for_start FROM provided_houses WHERE id = ?",
                                (house_id,),
                            ).fetchone()
                            if not payload["ready"]:
                                _clear_house_survey_records(db, house_id)
                            db.execute(
                                "UPDATE provided_houses SET ready_for_start = ? WHERE id = ?",
                                (int(payload["ready"]), house_id),
                            )
                            if bool(current["ready_for_start"]) != payload["ready"]:
                                _record_house_status(
                                    db,
                                    house_id,
                                    "ready" if payload["ready"] else "not_started",
                                    user["display_name"],
                                )
                self._send_json(200, {"id": house_id, "saved": True})
            elif path == "/api/stocktakes":
                payload = self._json_body()
                material_id = payload.get("material_id")
                location = str(payload.get("location") or "").strip()
                quantity = parse_quantity(payload.get("quantity"), "Counted quantity", True)
                counted_by = str(payload.get("counted_by") or "").strip()
                notes = str(payload.get("notes") or "").strip()
                if len(notes) > 2000:
                    raise ApiError(400, "Notes must be 2,000 characters or fewer.")
                with connect_db() as db:
                    with db:
                        db.execute("BEGIN IMMEDIATE")
                        if not db.execute(
                            "SELECT 1 FROM materials WHERE id = ?", (material_id,)
                        ).fetchone():
                            raise ApiError(400, "Choose a valid material.")
                        if location == "consumed" or location not in known_locations(db):
                            raise ApiError(400, "Choose a warehouse, team, or house with a recorded balance.")
                        _valid_reference(db, location)
                        if not counted_by:
                            raise ApiError(400, "Record who completed the stocktake.")
                        movement_id = db.execute(
                            "SELECT COALESCE(MAX(id), 0) FROM movements"
                        ).fetchone()[0]
                        db.execute(
                            """
                            INSERT INTO stocktakes
                                (material_id, location, quantity, movement_id, counted_on, counted_by, notes)
                            VALUES (?, ?, ?, ?, ?, ?, ?)
                            """,
                            (
                                material_id,
                                location,
                                quantity,
                                movement_id,
                                date.today().isoformat(),
                                counted_by,
                                notes,
                            ),
                        )
                self._send_json(201, {"ok": True})
            elif path == "/api/materials":
                payload = self._json_body()
                name = str(payload.get("name") or "").strip()
                code = str(payload.get("code") or "").strip()
                unit = str(payload.get("unit") or "units").strip() or "units"
                opening_qty = parse_quantity(
                    payload.get("opening_qty"), "Opening stock", allow_zero=True
                )
                if not name or len(name) > 200:
                    raise ApiError(400, "Material name is required (maximum 200 characters).")
                if len(code) > 100 or len(unit) > 40:
                    raise ApiError(400, "Material code or unit is too long.")
                with connect_db() as db:
                    with db:
                        cursor = db.execute(
                            "INSERT INTO materials(code, name, unit, opening_qty) VALUES (?, ?, ?, ?)",
                            (code, name, unit, opening_qty),
                        )
                self._send_json(201, {"id": cursor.lastrowid})
            elif path == "/api/calculator/requests":
                user = self._login_required()
                if user is None:
                    return
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "Request data must be a JSON object.")
                house_id = str(payload.get("house_id") or "").strip()
                data = payload.get("data")
                if not isinstance(data, dict):
                    raise ApiError(400, "Material request details are required.")
                encoded = json.dumps(data, ensure_ascii=False)
                if len(encoded) > 100_000:
                    raise ApiError(400, "Material request is too large.")
                notes = str(payload.get("notes") or "").strip()[:2000]
                with connect_db() as db:
                    if not db.execute(
                        "SELECT 1 FROM provided_houses WHERE id = ?", (house_id,)
                    ).fetchone():
                        raise ApiError(404, "Select a house from the provided list first.")
                    if db.execute(
                        "SELECT 1 FROM houses WHERE id = ? AND is_finished = 1",
                        (house_id,),
                    ).fetchone():
                        raise ApiError(409, f"House {house_id} is closed. Reopen it before assigning materials.")
                    cursor = db.execute(
                        """
                        INSERT INTO material_requests(house_id, created_at, created_by, notes, data)
                        VALUES (?, ?, ?, ?, ?)
                        """,
                        (
                            house_id,
                            datetime.now().isoformat(timespec="seconds"),
                            user["display_name"],
                            notes,
                            encoded,
                        ),
                    )
                self._send_json(201, {"id": cursor.lastrowid})
            elif path == "/api/houses":
                payload = self._json_body()
                if not isinstance(payload, dict):
                    raise ApiError(400, "House data must be a JSON object.")
                house_id = str(payload.get("id") or "").strip()
                if len(house_id) > 120:
                    raise ApiError(400, "House ID must be 120 characters or fewer.")
                postcode = str(payload.get("postcode") or "").strip()
                owner_name = str(payload.get("owner_name") or "").strip()
                employee_name = str(payload.get("employee_name") or "").strip()
                if len(postcode) > 80 or len(owner_name) > 160 or len(employee_name) > 160:
                    raise ApiError(400, "House details exceed the allowed length.")
                with connect_db() as db:
                    with db:
                        _require_ready_to_activate(db, house_id)
                        if not employee_name:
                            raise ApiError(
                                400,
                                "Assign an employee before opening the calculator.",
                            )
                        try:
                            perimeter = float(payload.get("perimeter_m"))
                        except (TypeError, ValueError, OverflowError):
                            raise ApiError(
                                400,
                                "Enter a surveyed perimeter greater than 0 and no more than 10,000 metres.",
                            )
                        if not math.isfinite(perimeter) or not (0 < perimeter <= 10000):
                            raise ApiError(
                                400,
                                "Enter a surveyed perimeter greater than 0 and no more than 10,000 metres.",
                            )
                        survey_row = db.execute(
                            "SELECT data FROM house_surveys WHERE house_id = ?",
                            (house_id,),
                        ).fetchone()
                        try:
                            survey_data = json.loads(survey_row["data"])
                        except (TypeError, json.JSONDecodeError):
                            raise ApiError(500, "Stored house survey data is invalid.")
                        if not isinstance(survey_data, dict):
                            raise ApiError(500, "Stored house survey data is invalid.")
                        survey_data["perimeter_m"] = perimeter
                        db.execute(
                            "UPDATE house_surveys SET data = ?, updated_by = ?, "
                            "updated_at = ? WHERE house_id = ?",
                            (
                                json.dumps(survey_data, ensure_ascii=False),
                                user["display_name"],
                                datetime.now().isoformat(timespec="seconds"),
                                house_id,
                            ),
                        )
                        was_active = db.execute(
                            "SELECT 1 FROM houses WHERE id = ?", (house_id,)
                        ).fetchone() is not None
                        provided_house = db.execute(
                            "SELECT postcode, owner_name FROM provided_houses WHERE id = ?",
                            (house_id,),
                        ).fetchone()
                        if provided_house:
                            postcode = provided_house["postcode"] or postcode
                            owner_name = provided_house["owner_name"] or owner_name
                        _upsert_house(
                            db,
                            house_id,
                            postcode,
                            owner_name,
                            employee_name,
                        )
                        if not was_active:
                            _record_house_status(
                                db, house_id, "active", user["display_name"]
                            )
                self._send_json(201, {"id": house_id})
            elif path == "/api/houses/status":
                payload = self._json_body()
                if not isinstance(payload, dict) or not isinstance(payload.get("finished"), bool):
                    raise ApiError(400, "Choose whether the house is finished.")
                house_id = str(payload.get("id") or "").strip()
                if not house_id:
                    raise ApiError(400, "House ID is required.")
                closeout_count = 0
                restored_count = 0
                with connect_db() as db:
                    with db:
                        db.execute("BEGIN IMMEDIATE")
                        house = db.execute(
                            "SELECT id, employee_name, is_finished FROM houses WHERE id = ?",
                            (house_id,),
                        ).fetchone()
                        if house is None:
                            raise ApiError(404, "House was not found.")
                        if not payload["finished"]:
                            restored_count = _undo_auto_closeouts(db, house_id)
                        if payload["finished"]:
                            material_rows = db.execute(
                                "SELECT id, name FROM materials ORDER BY name COLLATE NOCASE"
                            ).fetchall()
                            house_location = f"house:{house_id}"
                            balances = [
                                (material, balances_for(db, material["id"]).get(house_location, 0))
                                for material in material_rows
                            ]
                            negative = [
                                (material["name"], balance)
                                for material, balance in balances
                                if balance < -1e-9
                            ]
                            if negative:
                                detail = ", ".join(
                                    f"{name}: {balance:g}" for name, balance in negative
                                )
                                raise ApiError(
                                    409,
                                    f"Cannot finish this house while material balances are negative. Reconcile first: {detail}.",
                                )
                            for material, balance in balances:
                                if balance <= 1e-9:
                                    continue
                                _record_movement(
                                    db,
                                    {
                                        "kind": "use",
                                        "date": date.today().isoformat(),
                                        "material_id": material["id"],
                                        "qty": balance,
                                        "house_id": house_id,
                                        "person": user["display_name"],
                                        "team": "",
                                        "notes": CLOSEOUT_NOTE,
                                    },
                                )
                                closeout_count += 1
                        cursor = db.execute(
                            """
                            UPDATE houses
                            SET is_finished = ?
                            WHERE id = ?
                            """,
                            (
                                int(payload["finished"]),
                                house_id,
                            ),
                        )
                        if bool(house["is_finished"]) != payload["finished"]:
                            _record_house_status(
                                db,
                                house_id,
                                "closed" if payload["finished"] else "active",
                                user["display_name"],
                            )
                self._send_json(
                    200,
                    {
                        "id": house_id,
                        "finished": payload["finished"],
                        "materials_closed_out": closeout_count,
                        "materials_restored": restored_count,
                    },
                )
            elif path == "/api/houses/import":
                user = self._login_required()
                if user is None:
                    return
                if not self._admin_required(user):
                    return
                content_type = self.headers.get("Content-Type", "")
                if not content_type.startswith("multipart/form-data;"):
                    raise ApiError(400, "Upload the approved house-location CSV file.")
                body = self._body()
                match = re.search(r'boundary="?([^";]+)"?', content_type)
                if not match:
                    raise ApiError(400, "Invalid house-location file upload.")
                boundary = match.group(1).encode("ascii")
                file_data = None
                for part in body.split(b"--" + boundary):
                    if b"filename=" not in part or b"\r\n\r\n" not in part:
                        continue
                    headers, candidate = part.split(b"\r\n\r\n", 1)
                    filename = re.search(rb'filename="([^"]*)"', headers)
                    if filename and not filename.group(1).lower().endswith(b".csv"):
                        raise ApiError(400, "House locations must be uploaded as a CSV file.")
                    file_data = candidate.rsplit(b"\r\n", 1)[0]
                    break
                if not file_data:
                    raise ApiError(400, "Choose a house-location CSV file to import.")
                with connect_db() as db:
                    with db:
                        result = _import_house_locations_csv(db, file_data)
                self._send_json(201, result)
            elif path == "/api/import":
                content_type = self.headers.get("Content-Type", "")
                if not content_type.startswith("multipart/form-data;"):
                    raise ApiError(400, "Upload an Excel workbook using the file picker.")
                body = self._body()
                match = re.search(r'boundary="?([^";]+)"?', content_type)
                if not match:
                    raise ApiError(400, "Invalid file upload.")
                boundary = match.group(1).encode("ascii")
                parts = body.split(b"--" + boundary)
                file_data = None
                for part in parts:
                    if b"filename=" in part and b"\r\n\r\n" in part:
                        file_data = part.split(b"\r\n\r\n", 1)[1].rsplit(b"\r\n", 1)[0]
                        break
                if not file_data:
                    raise ApiError(400, "Choose an Excel workbook to import.")
                temporary_file = ROOT / "instance" / "workbook-upload.xlsx"
                temporary_file.parent.mkdir(parents=True, exist_ok=True)
                try:
                    temporary_file.write_bytes(file_data)
                    with connect_db() as db:
                        with db:
                            result = _import_workbook(db, temporary_file)
                    self._send_json(201, result)
                finally:
                    temporary_file.unlink(missing_ok=True)
            else:
                self._send_json(404, {"error": "Page not found."})
        except ApiError as exc:
            self._send_json(exc.status, {"error": str(exc)})
        except sqlite3.IntegrityError as exc:
            error_text = str(exc)
            if "materials.name" in error_text:
                message = "A material with that name already exists."
            elif "app_users.username" in error_text:
                message = "That username is already in use."
            else:
                message = "This record conflicts with existing tracker data."
            self._send_json(409, {"error": message})
        except (OSError, sqlite3.Error) as exc:
            self._send_json(500, {"error": f"Could not save the record: {exc}"})


def run():
    initialize_db()
    if bool(TLS_CERT) != bool(TLS_KEY):
        raise RuntimeError("Set both EWI_TLS_CERT and EWI_TLS_KEY to enable HTTPS.")
    if HOST not in {"127.0.0.1", "localhost", "::1"} and not COOKIE_SECURE:
        raise RuntimeError(
            "Binding to a network interface requires HTTPS; configure EWI_TLS_CERT and EWI_TLS_KEY."
        )
    with connect_db() as db:
        if HOST not in {"127.0.0.1", "localhost", "::1"} and not _user_count(db):
            raise RuntimeError(
                "Create the first administrator on localhost before enabling network access."
            )
    server = ThreadingHTTPServer((HOST, PORT), RequestHandler)
    scheme = "https" if COOKIE_SECURE else "http"
    if COOKIE_SECURE:
        tls_context = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
        tls_context.minimum_version = ssl.TLSVersion.TLSv1_2
        tls_context.load_cert_chain(TLS_CERT, TLS_KEY)
        server.socket = tls_context.wrap_socket(server.socket, server_side=True)
    print(f"EWI Tracker running at {scheme}://{HOST}:{PORT}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopping EWI Tracker.")
    finally:
        server.server_close()


if __name__ == "__main__":
    run()
